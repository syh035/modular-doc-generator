"""示例模板生成（UI「模板制作指南」页下载用）。

python-docx 现生成，不落盘不入 git；显式声明 eastAsia 中文字体
（M4/P17 教训：无 rFonts eastAsia 时 LO 会把中文渲染成空白）。
"""

import io

from docx import Document
from docx.document import Document as DocumentObj
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.run import Run

# 本机未装该字体也无妨：LO 解析到 eastAsia 声明即替换到系统 CJK 字体（M4 实锤）
_CJK_FONT = "宋体"


def _set_cjk_font(run: Run, *, size: int | None = None, bold: bool = False) -> None:
    """run 显式声明中文字体：font.name（ascii/hAnsi）+ rFonts eastAsia 双设置。"""
    run.font.name = _CJK_FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), _CJK_FONT)  # noqa: SLF001
    if size is not None:
        run.font.size = Pt(size)
    run.font.bold = bold


def _add_para(doc: DocumentObj, text: str, *, size: int | None = None, bold: bool = False) -> None:
    run = doc.add_paragraph().add_run(text)
    _set_cjk_font(run, size=size, bold=bold)


def build_sample_docx() -> bytes:
    """生成示例简历模板 docx（占位符覆盖：单段多占位符/多行块/成段正文）。"""
    doc = Document()
    _add_para(doc, "{{姓名}} 的简历", size=16, bold=True)
    _add_para(doc, "姓名：{{姓名}}    手机号：{{手机号}}")
    _add_para(doc, "教育经历", bold=True)
    _add_para(doc, "{{教育经历}}")
    _add_para(doc, "工作经历", bold=True)
    _add_para(doc, "{{工作经历}}")
    _add_para(doc, "自我评价", bold=True)
    _add_para(doc, "{{自我评价}}")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
