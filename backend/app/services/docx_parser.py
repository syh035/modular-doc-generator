"""占位符解析（D3 主路径）：从 DOCX 正文提取 `{{字段名}}` 可替换区域。

陷阱与决策落地：
- P1：Word 常把 `{{姓名}}` 拆成多个 run（`{{姓` + `名}}`）——必须段落内
  合并全部 run 文本后再匹配，禁止按 run 直读
- D6：v1 只扫 body 段落与表格单元格（含嵌套表格）；页眉页脚在独立
  header/footer part 中、文本框文字藏在 w:r/w:pict 深处——只取 w:p
  直接子 w:r 的 w:t 文本，两者天然落在解析范围外
- P3：区域身份锚定文档流元素——anchor.path 记录从 body 出发的定位
  路径（重放规则见 parse_placeholders 文档），order_index 为文档流顺序
"""

import re
from collections.abc import Iterator
from dataclasses import dataclass
from io import BytesIO
from typing import cast

from docx import Document
from lxml import etree

from app.services.field_lexicon import (
    CONF_PARAGRAPH,
    CONF_PLACEHOLDER,
    PARAGRAPH_MIN_LEN,
    classify_field,
    classify_paragraph,
    truncate_label,
)

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# `{{字段名}}`：容忍字段名首尾空白；字段名内不允许花括号（防嵌套误吞）
_PLACEHOLDER_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


@dataclass(frozen=True, slots=True)
class ParsedRegion:
    """占位符解析产物（type 恒为 custom，字段名词表启发式属 M3b）。"""

    placeholder: str  # 占位符原文，如 "{{姓名}}"
    label: str  # 显示名（字段名），如 "姓名"
    anchor: dict[str, object]  # 文档流锚点：{"kind": "p"|"cell_p", "path": [...]}
    order_index: int  # 文档流出现顺序（D7 迁移匹配依据）


@dataclass(frozen=True, slots=True)
class ParsedCandidate:
    """候选可替换区域（M3b 三层识别：占位符 > 字段名 > 成段正文）。"""

    region_type: str
    label: str  # 显示名：字段名 / 词表项 / 成段正文截断前缀
    placeholder: str | None  # 占位符区域才有；词表/成段候选为 None（整段替换）
    anchor: dict[str, object]
    order_index: int
    confidence: float


@dataclass(frozen=True, slots=True)
class ParseOutcome:
    """候选解析总产出：has_any_text 供纯图片模板判定（M3b）。"""

    candidates: list[ParsedCandidate]
    has_any_text: bool


def _paragraph_text(p: etree._Element) -> str:
    # 只拼直接子 w:r 的 w:t 文本（P1 合并 run；D6 天然排除文本框内文字）
    return "".join(t.text or "" for t in p.findall(f"{_W}r/{_W}t"))


_FlowElement = tuple[dict[str, object], etree._Element]


def iter_flow_paragraph_elements(body: etree._Element) -> Iterator[_FlowElement]:
    """按文档流顺序枚举段落元素引用（含表格单元格与嵌套表格）。

    anchor 编码与 iter_flow_paragraphs 完全一致——M6a 替换引擎以元素
    引用定位目标段落：lxml 引用在后续插入兄弟节点后依然有效，多区域
    替换（含多行克隆插入）互不干扰；替换完成后重新枚举即得新 path。
    """

    def scan_table(tbl: etree._Element, path: list[int]) -> Iterator[_FlowElement]:
        for row_idx, tr in enumerate(tbl.findall(f"{_W}tr")):
            for cell_idx, tc in enumerate(tr.findall(f"{_W}tc")):
                # 单元格内段落与嵌套表格各自独立计数（path 重放的依据）
                para_idx = 0
                nested_idx = 0
                for child in tc.iterchildren():
                    if child.tag == f"{_W}p":
                        yield (
                            {
                                "kind": "cell_p",
                                "path": [*path, row_idx, cell_idx, para_idx],
                            },
                            child,
                        )
                        para_idx += 1
                    elif child.tag == f"{_W}tbl":
                        yield from scan_table(child, [*path, row_idx, cell_idx, nested_idx])
                        nested_idx += 1

    for block_idx, child in enumerate(body.iterchildren()):
        if child.tag == f"{_W}p":
            yield {"kind": "p", "path": [block_idx]}, child
        elif child.tag == f"{_W}tbl":
            yield from scan_table(child, [block_idx])


def iter_flow_paragraphs(data: bytes) -> Iterator[tuple[dict[str, object], str]]:
    """按文档流顺序枚举 DOCX 全部正文段落（含表格单元格与嵌套表格）。

    每项 yield (anchor, 合并 run 后的段落全文)。anchor 编码与
    parse_placeholders 完全一致（P3，重放规则见其文档字符串）——
    M4 渲染位置匹配以此遍历为单一事实源，保证锚点与解析同源。
    """
    body = Document(BytesIO(data)).element.body
    for anchor, p in iter_flow_paragraph_elements(body):
        yield anchor, _paragraph_text(p)


def parse_placeholders(data: bytes) -> list[ParsedRegion]:
    """解析 DOCX bytes，按文档流顺序返回全部占位符区域。

    anchor.path 编码（P3，供后续模块重放定位元素）：
    - 段落：[block] —— body 第 block 个直接子元素即该 w:p
    - 单元格段落：[tbl, row, cell, para] —— body 第 tbl 个直接子元素是
      w:tbl → 第 row 个 w:tr → 第 cell 个 w:tc → 格内第 para 个直接子 w:p
    - 嵌套表格：每深入一层表格追加三元组 (tbl, row, cell)，末位仍是 para
    """
    regions: list[ParsedRegion] = []
    order = 0
    for anchor, text in iter_flow_paragraphs(data):
        for m in _PLACEHOLDER_RE.finditer(text):
            regions.append(
                ParsedRegion(
                    placeholder=m.group(0),
                    label=m.group(1),
                    anchor={
                        "kind": anchor["kind"],
                        "path": list(cast("list[int]", anchor["path"])),
                    },
                    order_index=order,
                )
            )
            order += 1
    return regions


def parse_candidates(data: bytes) -> ParseOutcome:
    """解析 DOCX，按文档流顺序产出全部候选区域（M3b 三层识别，PRD 4.2）。

    - 占位符区域：同 parse_placeholders，confidence=1.0；label 过词表
      推断 type（决策 C，如 {{姓名}} → name），未命中为 custom
    - 字段名区域：无占位符段落匹配词表（标题行 0.9 / 「字段名：值」0.7），
      placeholder=None（替换走整段替换，replacement._replace_whole_paragraph）
    - 成段正文：长度 ≥ PARAGRAPH_MIN_LEN 的无占位符段 → custom / 0.4
    - 优先级：有占位符的段落不再产出段落级候选（一段不重复建区域）；
      空段与过短非词表段（如「张三」）不产候选，但计入 has_any_text
    """
    candidates: list[ParsedCandidate] = []
    order = 0
    has_any_text = False
    for anchor, text in iter_flow_paragraphs(data):
        if text.strip():
            has_any_text = True
        ph_matches = list(_PLACEHOLDER_RE.finditer(text))
        if ph_matches:
            for m in ph_matches:
                label = m.group(1)
                candidates.append(
                    ParsedCandidate(
                        region_type=classify_field(label) or "custom",
                        label=label,
                        placeholder=m.group(0),
                        anchor={
                            "kind": anchor["kind"],
                            "path": list(cast("list[int]", anchor["path"])),
                        },
                        order_index=order,
                        confidence=CONF_PLACEHOLDER,
                    )
                )
                order += 1
            continue  # 占位符段不再产出段落级候选
        stripped = text.strip()
        if not stripped:
            continue
        hit = classify_paragraph(stripped)
        if hit is not None:
            candidates.append(
                ParsedCandidate(
                    region_type=hit.region_type,
                    label=stripped if hit.is_title else hit.term,
                    placeholder=None,
                    anchor={
                        "kind": anchor["kind"],
                        "path": list(cast("list[int]", anchor["path"])),
                    },
                    order_index=order,
                    confidence=hit.confidence,
                )
            )
            order += 1
        elif len(stripped) >= PARAGRAPH_MIN_LEN:
            candidates.append(
                ParsedCandidate(
                    region_type="custom",
                    label=truncate_label(stripped),
                    placeholder=None,
                    anchor={
                        "kind": anchor["kind"],
                        "path": list(cast("list[int]", anchor["path"])),
                    },
                    order_index=order,
                    confidence=CONF_PARAGRAPH,
                )
            )
            order += 1
    return ParseOutcome(candidates, has_any_text)
