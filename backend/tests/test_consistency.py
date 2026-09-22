"""一致性专项测试（里程碑 4）：diff 工具单测 + 五类模板端到端（预览 vs 导出重转）。

端到端口径（TODO 里程碑 4）：单管线铁律 + P24 确定性序列化 → 导出 DOCX
与渲染产物同字节 → LO 内容寻址缓存命中同一 PDF → 预览 PDF 与导出重转
PDF 字节级一致（强口径）；内容级逐页容差 diff（services/consistency.py）
作兜底断言。五类模板 = 纯文本单栏 / 带表格 / 双栏 / 带文本框 / 带占位符。
"""

import hashlib
from io import BytesIO
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

from app.services.consistency import diff_pdfs
from app.services.libreoffice import LibreOfficeManager, find_soffice, get_manager
from app.services.render_service import render_version

_EAST = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia"

soffice_available = pytest.mark.skipif(find_soffice() is None, reason="LibreOffice 未安装")


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _cjk_run(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    run = p.runs[0]
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(_EAST, "宋体")


def _cell_run(cell: object, text: str) -> None:
    p = cell.paragraphs[0]  # type: ignore[attr-defined]
    run = p.add_run(text)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(_EAST, "宋体")


def _add_textbox(doc: Document, text: str) -> None:
    """VML 文本框（D6：解析跳过、导出保留原样）。"""
    p = doc.add_paragraph()
    run = p.add_run()
    pict = parse_xml(
        '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        ' xmlns:v="urn:schemas-microsoft-com:vml">'
        '<v:shape style="width:200pt;height:60pt">'
        "<v:textbox><w:txbxContent>"
        f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox>"
        "</v:shape>"
        "</w:pict>"
    )
    run._r.append(pict)


def _set_two_columns(doc: Document) -> None:
    """当前节设双栏（CT_SectPr 无 get_or_add_cols，直接操作 w:cols）。"""
    sect_pr = doc.sections[0]._sectPr
    cols = sect_pr.xpath("./w:cols")
    if cols:
        cols[0].set(qn("w:num"), "2")
    else:
        cols_el = OxmlElement("w:cols")
        cols_el.set(qn("w:num"), "2")
        sect_pr.append(cols_el)


def _docx_bytes(doc: Document) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _convert(lo: LibreOfficeManager, doc: Document, name: str, tmp_path: Path) -> Path:
    path = tmp_path / name
    path.write_bytes(_docx_bytes(doc))
    return lo.convert(path)


# ---- diff 工具单测（真实 LO 产物）----


@pytest.fixture
def lo(tmp_path: Path) -> LibreOfficeManager:
    soffice = find_soffice()
    assert soffice is not None
    return LibreOfficeManager(
        str(soffice),
        profile_dir=tmp_path / "prof",
        cache_dir=tmp_path / "cache",
        timeout_seconds=120.0,
    )


@soffice_available
def test_diff_identical_pdf(tmp_path: Path, lo: LibreOfficeManager) -> None:
    doc = Document()
    _cjk_run(doc, "第一行测试文本")
    _cjk_run(doc, "第二行测试文本")
    pdf = _convert(lo, doc, "same.docx", tmp_path)
    report = diff_pdfs(pdf, pdf)
    assert report.consistent
    assert report.pages_equal and report.page_count_a == 1 and report.diffs == []


@soffice_available
def test_diff_text_change_reports_missing_and_extra(tmp_path: Path, lo: LibreOfficeManager) -> None:
    doc_a = Document()
    _cjk_run(doc_a, "旧文本内容行")
    pdf_a = _convert(lo, doc_a, "a.docx", tmp_path)
    doc_b = Document()
    _cjk_run(doc_b, "新文本内容行")
    pdf_b = _convert(lo, doc_b, "b.docx", tmp_path)

    report = diff_pdfs(pdf_a, pdf_b)
    assert not report.consistent
    kinds = {(d.kind, d.text) for d in report.diffs}
    assert ("missing", "旧文本内容行") in kinds
    assert ("extra", "新文本内容行") in kinds


@soffice_available
def test_diff_moved_beyond_tol_and_absorbed_by_large_tol(
    tmp_path: Path, lo: LibreOfficeManager
) -> None:
    doc_a = Document()
    _cjk_run(doc_a, "位移测试行")
    pdf_a = _convert(lo, doc_a, "a.docx", tmp_path)
    doc_b = Document()
    for _ in range(20):
        doc_b.add_paragraph("")
    _cjk_run(doc_b, "位移测试行")
    pdf_b = _convert(lo, doc_b, "b.docx", tmp_path)

    report = diff_pdfs(pdf_a, pdf_b)
    moved = [d for d in report.diffs if d.kind == "moved"]
    assert len(moved) == 1
    assert moved[0].text == "位移测试行" and moved[0].delta > 100
    assert moved[0].bbox_a is not None and moved[0].bbox_b is not None
    # 容差放大到位移以内 → 一致
    assert diff_pdfs(pdf_a, pdf_b, coord_tol=1000.0).consistent


@soffice_available
def test_diff_page_count_mismatch_reports_extra_page(
    tmp_path: Path, lo: LibreOfficeManager
) -> None:
    doc_a = Document()
    _cjk_run(doc_a, "内容行000")
    pdf_a = _convert(lo, doc_a, "a.docx", tmp_path)
    doc_b = Document()
    for i in range(120):
        _cjk_run(doc_b, f"内容行{i:03d}")
    pdf_b = _convert(lo, doc_b, "b.docx", tmp_path)

    report = diff_pdfs(pdf_a, pdf_b)
    assert report.pages_equal is False
    assert report.page_count_a == 1 and report.page_count_b >= 2
    # B 多出的页整页记 extra（0 基页码 ≥ 1）
    assert {d.page for d in report.diffs if d.kind == "extra"} >= {1}
    assert all(d.kind != "missing" for d in report.diffs)


# ---- 五类模板端到端：预览 PDF vs 导出重转 PDF ----


def _upload(client: TestClient, data: bytes) -> dict:
    resp = client.post(
        "/api/templates",
        files={"file": ("模板.docx", data, "application/octet-stream")},
    )
    assert resp.status_code == 201
    return resp.json()


def _make_block(client: TestClient, name: str, content: str) -> int:
    resp = client.post("/api/blocks", json={"name": name, "content": content})
    assert resp.status_code == 201
    return resp.json()["id"]


def _region_by_placeholder(regions: list[dict], placeholder: str) -> dict:
    matches = [r for r in regions if r["placeholder"] == placeholder]
    assert len(matches) == 1, f"占位符 {placeholder} 区域数异常：{len(matches)}"
    return matches[0]


def _bind(client: TestClient, version_id: int, region_id: int, block_id: int) -> None:
    resp = client.post(
        f"/api/versions/{version_id}/bindings",
        json={"region_id": region_id, "block_id": block_id},
    )
    assert resp.status_code == 200


def _bind_ph(
    client: TestClient,
    version_id: int,
    regions: list[dict],
    placeholder: str,
    block_id: int,
) -> None:
    _bind(client, version_id, _region_by_placeholder(regions, placeholder)["id"], block_id)


def _assert_export_consistent(
    client: TestClient, tmp_path: Path, version_id: int, confirm: bool = False
) -> Path:
    """走完整链：预览 → 导出落盘 → LO 重转，断言字节一致 + 容差 diff 零差异。"""
    preview = client.get(f"/api/versions/{version_id}/preview")
    assert preview.status_code == 200
    preview_path = tmp_path / "preview.pdf"
    preview_path.write_bytes(preview.content)

    exported = client.post(
        f"/api/versions/{version_id}/export",
        json={"confirm_large_overflow": confirm},
    )
    assert exported.status_code == 200, exported.json()
    exported_path = tmp_path / "exported.docx"
    exported_path.write_bytes(exported.content)

    # 强口径①：导出 DOCX 与渲染产物同字节（单管线铁律 + P24 确定性序列化）
    assert _sha(render_version(version_id).data) == _sha(exported.content)

    # 强口径②：重转按内容 sha 命中缓存 → 与预览 PDF 字节一致
    reconverted = get_manager().convert(exported_path)
    assert _sha(reconverted.read_bytes()) == _sha(preview.content)

    # 兜底口径：内容级逐页容差 diff 零差异
    report = diff_pdfs(preview_path, reconverted)
    assert report.consistent, report.to_dict()
    return reconverted


@soffice_available
def test_e2e_plain_text_paragraph(client: TestClient, tmp_path: Path) -> None:
    """类 1 纯文本单栏：成段正文区域（无占位符）整段替换路径。"""
    doc = Document()
    _cjk_run(doc, "张三，男，某大学计算机科学与技术专业毕业，求职意向为数据工程师岗位方向。")
    _cjk_run(doc, "负责数据平台建设与维护，主导三个核心系统的架构设计落地，持续优化任务调度。")
    tpl = _upload(client, _docx_bytes(doc))
    assert len(tpl["regions"]) == 2
    assert all(r["placeholder"] is None for r in tpl["regions"])

    vid = tpl["default_version_id"]
    intro = _make_block(client, "简介块", "李四，数据平台工程师。")
    _bind(client, vid, tpl["regions"][0]["id"], intro)
    _assert_export_consistent(client, tmp_path, vid)


@soffice_available
def test_e2e_table_placeholder(client: TestClient, tmp_path: Path) -> None:
    """类 2 带表格：单元格占位符替换（词表命中教育背景）。"""
    doc = Document()
    _cjk_run(doc, "张三的简历文档示例标题行")
    table = doc.add_table(rows=2, cols=2)
    _cell_run(table.cell(0, 0), "{{姓名}}")
    _cell_run(table.cell(0, 1), "1995 年 3 月")
    _cell_run(table.cell(1, 0), "{{教育背景}}")
    _cell_run(table.cell(1, 1), "某大学计算机专业本科")
    tpl = _upload(client, _docx_bytes(doc))
    assert len(tpl["regions"]) == 2

    vid = tpl["default_version_id"]
    name_block = _make_block(client, "姓名块", "李四")
    edu_block = _make_block(client, "教育块", "某交通大学硕士")
    _bind_ph(client, vid, tpl["regions"], "{{姓名}}", name_block)
    _bind_ph(client, vid, tpl["regions"], "{{教育背景}}", edu_block)
    _assert_export_consistent(client, tmp_path, vid)


@soffice_available
def test_e2e_two_columns(client: TestClient, tmp_path: Path) -> None:
    """类 3 双栏排版：两栏内占位符替换后渲染一致。"""
    doc = Document()
    _set_two_columns(doc)
    _cjk_run(doc, "{{联系方式}}")
    _cjk_run(doc, "电话一三八零零零零零零零零，邮箱测试账号，坐标上海市浦东新区软件园。")
    _cjk_run(doc, "{{技能特长}}")
    _cjk_run(doc, "熟练掌握数据管道开发与调优，具备大规模分布式集群运维实战经验积累。")
    tpl = _upload(client, _docx_bytes(doc))
    assert len(tpl["regions"]) == 4

    vid = tpl["default_version_id"]
    contact_block = _make_block(client, "联系块", "13800000000")
    skill_block = _make_block(client, "技能块", "Python")
    _bind_ph(client, vid, tpl["regions"], "{{联系方式}}", contact_block)
    _bind_ph(client, vid, tpl["regions"], "{{技能特长}}", skill_block)
    _assert_export_consistent(client, tmp_path, vid)


@soffice_available
def test_e2e_textbox_preserved(client: TestClient, tmp_path: Path) -> None:
    """类 4 带文本框：D6 文本框不解析不替换、导出原样保留；正文替换一致。"""
    doc = Document()
    _add_textbox(doc, "文本框装饰内容甲乙丙")
    _cjk_run(doc, "姓名：{{姓名}}")
    _cjk_run(doc, "工作经历描述段落需要足够长的文本以触发成段正文识别逻辑的阈值判断。")
    tpl = _upload(client, _docx_bytes(doc))
    assert len(tpl["regions"]) == 2

    vid = tpl["default_version_id"]
    name_block = _make_block(client, "姓名块", "李四")
    _bind_ph(client, vid, tpl["regions"], "{{姓名}}", name_block)
    # 成段正文区域（regions 中 placeholder=None）整段替换
    para_region = next(r for r in tpl["regions"] if r["placeholder"] is None)
    _bind(client, vid, para_region["id"], _make_block(client, "工作块", "负责数据平台建设。"))
    reconverted = _assert_export_consistent(client, tmp_path, vid)

    # D6 验证：文本框内容在重转 PDF 全文中保留（P18：文本框浮动文字按
    # 内容流乱序拆片，无法按序子串匹配，退化为字符全集断言）
    with pymupdf.open(reconverted) as pdf_doc:
        full_text = "".join(page.get_text() for page in pdf_doc)
    compact = full_text.replace("\n", "")
    assert all(ch in compact for ch in "文本框装饰内容甲乙丙")


@soffice_available
def test_e2e_multi_placeholder_with_multiline_block(client: TestClient, tmp_path: Path) -> None:
    """类 5 多占位符：同段双占位符 + 多行块克隆行（大超出确认导出 D5 路径）。"""
    doc = Document()
    _cjk_run(doc, "{{姓名}}")
    _cjk_run(doc, "{{联系方式}}与{{技能特长}}同段相邻")
    _cjk_run(doc, "工作内容：{{工作经历}}")
    tpl = _upload(client, _docx_bytes(doc))
    assert len(tpl["regions"]) == 4

    vid = tpl["default_version_id"]
    _bind_ph(client, vid, tpl["regions"], "{{姓名}}", _make_block(client, "姓名块", "李四"))
    _bind_ph(
        client, vid, tpl["regions"], "{{联系方式}}", _make_block(client, "联系块", "13800000000")
    )
    _bind_ph(client, vid, tpl["regions"], "{{技能特长}}", _make_block(client, "技能块", "Python"))
    # 多行块 → 克隆行 → 大超出 → 走 D5 确认导出
    work = _make_block(client, "工作块", "负责核心模块开发\n性能优化落地\n沉淀运维手册")
    _bind_ph(client, vid, tpl["regions"], "{{工作经历}}", work)
    _assert_export_consistent(client, tmp_path, vid, confirm=True)
