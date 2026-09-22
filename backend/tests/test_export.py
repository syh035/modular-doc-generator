"""导出测试（M9，PRD 4.8 / D5）：文件名/落盘/大超出确认/写盘失败 + LO 端到端。

单元部分 mock render_version（不触发 LO，快速确定性）；端到端部分走真实
渲染管线验证 409→确认→落盘全链路与产物可打开。
"""

from io import BytesIO
from pathlib import Path
from urllib.parse import unquote

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.core.config import settings
from app.services import render_service
from app.services.export_service import (
    export_file_name,
    extract_large_overflow_warnings,
    sanitize_filename_part,
)
from app.services.libreoffice import find_soffice

_EAST = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia"


# ---- 纯函数：文件名清洗与拼接 ----


def test_sanitize_filename_part() -> None:
    assert sanitize_filename_part("投递A岗") == "投递A岗"
    assert sanitize_filename_part("a/b\\c:d*e?f\"g<h>i|j") == "a_b_c_d_e_f_g_h_i_j"
    assert sanitize_filename_part("  .x. ") == "x"
    assert sanitize_filename_part("///") == "___"  # 非法字符换下划线后仍合法


def test_export_file_name() -> None:
    from datetime import datetime

    name = export_file_name("后端侧重版", datetime(2026, 9, 21, 10, 30))
    assert name == "简历-后端侧重版-20260921.docx"


def test_extract_warnings_only_large() -> None:
    def ov(level: str, ratio: float) -> dict[str, object]:
        return {"level": level, "ratio": ratio, "clipped": False, "fixed_row": False}

    items: list[dict[str, object]] = [
        {"id": 1, "label": "大", "overflow": ov("large", 2.89)},
        {"id": 2, "label": "小", "overflow": ov("small", 0.3)},
        {"id": 3, "label": "无", "overflow": None},
    ]
    warnings = extract_large_overflow_warnings(items)
    assert len(warnings) == 1
    assert warnings[0]["region_id"] == 1
    assert warnings[0]["label"] == "大"
    assert warnings[0]["ratio"] == 2.89


# ---- 端点单元（mock render_version，不依赖 LO）----


class _StubRender:
    """render_version 替身：data 交给真实落盘断言，items 控制警示路径。"""

    def __init__(self, data: bytes, items: list[dict[str, object]]) -> None:
        self.data = data
        self.pdf_path = Path("/tmp/unused.pdf")
        self.items = items


def _simple_docx() -> bytes:
    doc = Document()
    p = doc.add_paragraph("姓名：{{姓名}}")
    run = p.runs[0]
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(_EAST, "宋体")  # type: ignore[union-attr]
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _patch_render(
    monkeypatch: pytest.MonkeyPatch, data: bytes, items: list[dict[str, object]]
) -> None:
    monkeypatch.setattr(render_service, "render_version", lambda vid: _StubRender(data, items))


def _setup_version(client: TestClient, name: str | None = None) -> int:
    """上传模板建默认版本（不走渲染），可选重命名。"""
    resp = client.post(
        "/api/templates",
        files={"file": ("模板.docx", _simple_docx(), "application/octet-stream")},
    )
    assert resp.status_code == 201
    vid = resp.json()["default_version_id"]
    assert vid is not None
    if name is not None:
        assert client.patch(f"/api/versions/{vid}", json={"name": name}).status_code == 200
    return vid


def test_export_ok_writes_file_and_disposition(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    vid = _setup_version(client, "投递A岗")
    _patch_render(monkeypatch, _simple_docx(), [])

    resp = client.post(f"/api/versions/{vid}/export", json={})

    assert resp.status_code == 200
    assert resp.headers["cache-control"] == "no-store"
    assert "attachment" in resp.headers["content-disposition"]
    # 非 ASCII 文件名走 RFC 5987 filename*（starlette 行为），解码后校验规则
    encoded = resp.headers["content-disposition"].split("filename*=utf-8''", 1)[1]
    assert unquote(encoded).startswith("简历-投递A岗-")
    # 落盘存在且内容 = 同源产物；浏览器下载体 = 落盘文件
    files = list(settings.exports_dir.glob("简历-投递A岗-*.docx"))
    assert len(files) == 1
    assert files[0].read_bytes() == _simple_docx()
    assert resp.content == _simple_docx()


def test_export_document_opens_with_python_docx(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """落盘产物是合法 DOCX：python-docx 可打开且内容完整。"""
    vid = _setup_version(client)
    _patch_render(monkeypatch, _simple_docx(), [])

    assert client.post(f"/api/versions/{vid}/export", json={}).status_code == 200

    files = list(settings.exports_dir.glob("简历-默认版本-*.docx"))
    assert len(files) == 1
    doc = Document(str(files[0]))
    assert doc.paragraphs[0].text == "姓名：{{姓名}}"


def test_export_rerun_overwrites_same_name(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """同版本同日反复导出：同名覆盖，不累积残留文件。"""
    vid = _setup_version(client)
    _patch_render(monkeypatch, _simple_docx(), [])
    assert client.post(f"/api/versions/{vid}/export", json={}).status_code == 200
    assert client.post(f"/api/versions/{vid}/export", json={}).status_code == 200
    assert len(list(settings.exports_dir.glob("*.docx"))) == 1


def test_export_version_not_found(client: TestClient) -> None:
    resp = client.post("/api/versions/9999/export", json={})
    assert resp.status_code == 404
    assert resp.json()["error"]["code"] == "VERSION_NOT_FOUND"


def test_export_large_overflow_blocked_then_confirmed(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """D5 主路径：大超出未确认 409 + warnings 清单；确认后 200 落盘。"""
    vid = _setup_version(client, "有溢出版")
    large_items: list[dict[str, object]] = [
        {
            "id": 7,
            "label": "项目经历",
            "overflow": {"level": "large", "ratio": 1.5, "clipped": False, "fixed_row": False},
        }
    ]
    _patch_render(monkeypatch, _simple_docx(), large_items)

    blocked = client.post(f"/api/versions/{vid}/export", json={})
    assert blocked.status_code == 409
    body = blocked.json()
    assert body["error"]["code"] == "EXPORT_LARGE_OVERFLOW"
    assert body["warnings"] == [
        {"region_id": 7, "label": "项目经历", "ratio": 1.5, "clipped": False, "fixed_row": False}
    ]
    assert list(settings.exports_dir.glob("*.docx")) == []  # 拦截时不落盘

    confirmed = client.post(f"/api/versions/{vid}/export", json={"confirm_large_overflow": True})
    assert confirmed.status_code == 200
    assert len(list(settings.exports_dir.glob("简历-有溢出版-*.docx"))) == 1


def test_export_small_overflow_not_blocked(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """小超出不拦截：重排导出即默认行为，无需确认。"""
    vid = _setup_version(client)
    items: list[dict[str, object]] = [
        {
            "id": 7,
            "label": "自荐",
            "overflow": {"level": "small", "ratio": 0.3, "clipped": False, "fixed_row": False},
        }
    ]
    _patch_render(monkeypatch, _simple_docx(), items)
    assert client.post(f"/api/versions/{vid}/export", json={}).status_code == 200


def test_export_write_failure_readable(
    client: TestClient, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """写盘失败 → EXPORT_WRITE_FAILED 可读报错（非静默）。"""
    vid = _setup_version(client)
    _patch_render(monkeypatch, _simple_docx(), [])
    blocker = tmp_path / "exports_blocker"  # 目录位被同名文件占据 → mkdir 失败
    blocker.write_text("not a dir")
    monkeypatch.setattr(settings, "exports_dir", blocker)

    resp = client.post(f"/api/versions/{vid}/export", json={})
    assert resp.status_code == 500
    assert resp.json()["error"]["code"] == "EXPORT_WRITE_FAILED"


# ---- 端到端（真实 LO）：大超出版本全链路 ----


@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice 未安装")
def test_export_e2e_large_overflow_flow(client: TestClient) -> None:
    """真实管线：多行内容绑定 → 导出 409 警示 → 确认 → 落盘成品可打开。"""
    doc = Document()
    p = doc.add_paragraph("项目经历：{{项目经历}}")
    run = p.runs[0]
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(_EAST, "宋体")  # type: ignore[union-attr]
    buf = BytesIO()
    doc.save(buf)

    resp = client.post(
        "/api/templates",
        files={"file": ("e2e导出.docx", buf.getvalue(), "application/octet-stream")},
    )
    assert resp.status_code == 201
    body = resp.json()
    vid = body["default_version_id"]
    assert vid is not None
    rid = body["regions"][0]["id"]

    block = client.post(
        "/api/blocks",
        json={
            "name": "项目块",
            "content": (
                "独立负责数据迁移平台\n完成核心链路重构\n性能提升三倍\n沉淀运维手册\n输出季度复盘"
            ),
        },
    )
    assert block.status_code == 201
    assert (
        client.post(
            f"/api/versions/{vid}/bindings",
            json={"region_id": rid, "block_id": block.json()["id"]},
        ).status_code
        == 200
    )

    blocked = client.post(f"/api/versions/{vid}/export", json={})
    assert blocked.status_code == 409
    warnings = blocked.json()["warnings"]
    assert len(warnings) == 1
    assert warnings[0]["label"] == "项目经历"
    assert warnings[0]["clipped"] is False

    confirmed = client.post(f"/api/versions/{vid}/export", json={"confirm_large_overflow": True})
    assert confirmed.status_code == 200
    files = list(settings.exports_dir.glob("简历-默认版本-*.docx"))
    assert len(files) == 1
    exported = Document(str(files[0]))
    texts = "\n".join(p.text for p in exported.paragraphs)
    assert "独立负责数据迁移平台" in texts  # 绑定区域写入块内容
    assert "{{项目经历}}" not in texts
