"""替换引擎单测（M6a）：P1 跨 run / P4 编号剥离 / P5 样式继承 / 多行换段 /
path 推挤映射 / 未绑定保留 / 表格与嵌套表格。"""

import dataclasses
import json
from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.models.entities import Region
from app.services.docx_parser import parse_placeholders
from app.services.replacement import apply_replacements

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def docx_bytes(doc: Document) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def regions_of(data: bytes, start_id: int = 1) -> list[Region]:
    """按解析产物构造 Region 实体（anchor 落库形态：JSON 文本）。"""
    return [
        Region(
            id=start_id + i,
            template_id=1,
            type="custom",
            label=pr.label,
            placeholder=pr.placeholder,
            anchor=json.dumps(pr.anchor, ensure_ascii=False),
            order_index=pr.order_index,
            bbox_json=None,
            confidence=None,
            review_status="pending",
            created_at="",
            updated_at="",
        )
        for i, pr in enumerate(parse_placeholders(data))
    ]


def paragraphs_text(data: bytes) -> list[str]:
    """成品文档正文段落文本（顶层，按序）。"""
    return [p.text for p in Document(BytesIO(data)).paragraphs]


# ---- 基础替换 ----


def _doc_with(*texts: str) -> Document:
    doc = Document()
    for t in texts:
        doc.add_paragraph(t)
    return doc


def test_single_placeholder_single_line() -> None:
    data = docx_bytes(_doc_with("电话：{{手机号}}"))
    regions = regions_of(data)
    outcome = apply_replacements(data, regions, {regions[0].id: "13800138000"})
    assert paragraphs_text(outcome.data)[0] == "电话：13800138000"
    assert outcome.region_paths[regions[0].id] == [0]


def test_unbound_placeholder_kept_verbatim() -> None:
    """未绑定区域原样保留占位符（M9 导出同规则）。"""
    data = docx_bytes(_doc_with("{{姓名}}：{{电话}}"))
    regions = regions_of(data)
    outcome = apply_replacements(data, regions, {regions[0].id: "张三"})
    assert paragraphs_text(outcome.data)[0] == "张三：{{电话}}"
    # 未绑定区域同样给出 path（overlay 原位对齐用）
    assert outcome.region_paths[regions[1].id] == [0]


# ---- P1：占位符跨 run ----


def test_split_run_placeholder_replaced() -> None:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("联系{{手")
    p.add_run("机号}}感谢")
    data = docx_bytes(doc)
    regions = regions_of(data)
    assert len(regions) == 1

    outcome = apply_replacements(data, regions, {regions[0].id: "138"})
    assert paragraphs_text(outcome.data)[0] == "联系138感谢"


def test_split_run_style_inherits_first_covering_run() -> None:
    """P5：新文本由占位符首覆盖 run 承载（继承其 rPr，此处加粗）。"""
    doc = Document()
    p = doc.add_paragraph()
    r1 = p.add_run("前缀{{姓")
    r1.bold = True
    p.add_run("名}}后缀")
    data = docx_bytes(doc)
    regions = regions_of(data)

    outcome = apply_replacements(data, regions, {regions[0].id: "张三"})
    para = Document(BytesIO(outcome.data)).paragraphs[0]
    new_run = next(r for r in para.runs if "张三" in (r.text or ""))
    assert new_run.bold is True


# ---- D9：多行换段 + P4 编号剥离 ----


def test_multiline_content_clones_paragraphs() -> None:
    data = docx_bytes(_doc_with("经历：{{工作经历}}", "下一段"))
    regions = regions_of(data)
    content = "公司A\n岗位：后端\n业绩：xxx"
    outcome = apply_replacements(data, regions, {regions[0].id: content})

    texts = paragraphs_text(outcome.data)
    # 首行留原位，后续两行克隆段落插入，原「下一段」被推到最后
    assert texts == ["经历：公司A", "岗位：后端", "业绩：xxx", "下一段"]
    # path 映射：区域首行仍在 [0]；被推挤的后续段落不出现在映射（非区域）
    assert outcome.region_paths[regions[0].id] == [0]


def test_clone_paragraphs_strip_numbering() -> None:
    """P4：克隆段剥离编号属性，防 1. 2. 2. 3. 错乱。"""
    doc = Document()
    p = doc.add_paragraph("{{列表项}}")
    ppr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    for tag, val in (("w:ilvl", "0"), ("w:numId", "1")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        num_pr.append(el)
    ppr.append(num_pr)
    data = docx_bytes(doc)
    regions = regions_of(data)

    outcome = apply_replacements(data, regions, {regions[0].id: "一\n二"})
    paras = Document(BytesIO(outcome.data)).paragraphs
    assert [pp.text for pp in paras] == ["一", "二"]
    # P4 只管克隆段：原位段保留原 pPr（含编号），克隆段必须剥离
    assert paras[0]._p.find(f"{_W}pPr").find(f"{_W}numPr") is not None
    clone_ppr = paras[1]._p.find(f"{_W}pPr")
    assert clone_ppr is None or clone_ppr.find(f"{_W}numPr") is None


def test_region_paths_after_multiline_push() -> None:
    """多行插入推挤后续区域段落索引 → 映射以替换后重枚举为准。"""
    data = docx_bytes(_doc_with("{{姓名}}", "{{电话}}"))
    regions = regions_of(data)
    outcome = apply_replacements(data, regions, {regions[0].id: "张\n三\n丰"})
    # 姓名克隆 2 段 → 电话段落从 body[1] 推到 body[3]
    assert outcome.region_paths[regions[0].id] == [0]
    assert outcome.region_paths[regions[1].id] == [3]
    assert paragraphs_text(outcome.data) == ["张", "三", "丰", "{{电话}}"]


# ---- 同段多占位符 ----


def test_same_paragraph_mixed_binding() -> None:
    """同段两占位符：一绑一留 → 绑定者替换、未绑定者原文保留。"""
    data = docx_bytes(_doc_with("{{姓名}}·{{电话}}"))
    regions = regions_of(data)
    outcome = apply_replacements(data, regions, {regions[1].id: "138"})
    assert paragraphs_text(outcome.data)[0] == "{{姓名}}·138"


def test_same_paragraph_both_multiline_order_kept() -> None:
    """同段两占位符均多行：克隆行按区域顺序 × 行序插入。"""
    data = docx_bytes(_doc_with("A{{甲}}B{{乙}}C"))
    regions = regions_of(data)
    outcome = apply_replacements(
        data, regions, {regions[0].id: "甲1\n甲2", regions[1].id: "乙1\n乙2"}
    )
    assert paragraphs_text(outcome.data) == ["A甲1B乙1C", "甲2", "乙2"]


# ---- 表格（cell_p anchor）----


def test_table_cell_placeholder_replaced() -> None:
    doc = Document()
    doc.add_paragraph("标题{{姓名}}")
    tbl = doc.add_table(rows=1, cols=2)
    tbl.cell(0, 0).text = "{{电话}}"
    tbl.cell(0, 1).text = "固定内容"
    data = docx_bytes(doc)
    regions = regions_of(data)
    by_label = {r.label: r for r in regions}
    outcome = apply_replacements(
        data, regions, {by_label["电话"].id: "138", by_label["姓名"].id: "张三"}
    )
    doc2 = Document(BytesIO(outcome.data))
    assert doc2.paragraphs[0].text == "标题张三"
    assert doc2.tables[0].cell(0, 0).text == "138"
    assert doc2.tables[0].cell(0, 1).text == "固定内容"


def test_table_cell_multiline_clone_inside_cell() -> None:
    """表格内多行：克隆段插在单元格内（不逃出表格）。"""
    doc = Document()
    tbl = doc.add_table(rows=1, cols=1)
    tbl.cell(0, 0).text = "{{教育}}"
    data = docx_bytes(doc)
    regions = regions_of(data)
    outcome = apply_replacements(data, regions, {regions[0].id: "本科\n硕士"})
    cell = Document(BytesIO(outcome.data)).tables[0].cell(0, 0)
    assert [p.text for p in cell.paragraphs] == ["本科", "硕士"]


def test_nested_table_placeholder_replaced() -> None:
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    outer.cell(0, 0).text = "外层"
    inner = outer.cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "{{嵌套}}"
    data = docx_bytes(doc)
    regions = regions_of(data)
    assert len(regions) == 1
    outcome = apply_replacements(data, regions, {regions[0].id: "OK"})
    doc2 = Document(BytesIO(outcome.data))
    assert doc2.tables[0].cell(0, 0).tables[0].cell(0, 0).text == "OK"


# ---- 防御 ----


def test_stale_anchor_skipped_not_crash() -> None:
    """anchor 与文档失配（如手工改库）→ 区域跳过、不崩、不出映射。"""
    data = docx_bytes(_doc_with("{{姓名}}"))
    regions = regions_of(data)
    anchor = json.loads(regions[0].anchor)
    anchor["path"] = [99]
    stale = dataclasses.replace(regions[0], anchor=json.dumps(anchor))
    outcome = apply_replacements(data, [stale], {stale.id: "张三"})
    assert paragraphs_text(outcome.data)[0] == "{{姓名}}"
    assert outcome.region_paths == {}


def test_empty_content_replaces_to_empty() -> None:
    data = docx_bytes(_doc_with("A{{姓名}}B"))
    regions = regions_of(data)
    outcome = apply_replacements(data, regions, {regions[0].id: ""})
    assert paragraphs_text(outcome.data)[0] == "AB"
