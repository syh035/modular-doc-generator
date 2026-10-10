"""软回车保持同段、保留样式、复制设置且预览导出同源。"""

from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.models.db import get_conn
from app.models.repositories import bindings
from app.services.docx_parser import iter_flow_paragraphs
from app.services.libreoffice import find_soffice
from app.services.render_service import render_version


@pytest.fixture
def soft_scene(client: TestClient) -> tuple[dict, dict]:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("前缀{{内容}}后缀").bold = True
    doc.add_paragraph("结尾")
    b = BytesIO()
    doc.save(b)
    tpl = client.post("/api/templates", files={"file": ("软回车.docx", b.getvalue())}).json()
    block = client.post("/api/blocks", json={"name": "多行", "content": "第一行\n第二行"}).json()
    return tpl, block


@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice 未安装")
def test_soft_break_preview_export_and_version_copy(
    client: TestClient, soft_scene: tuple[dict, dict]
) -> None:
    tpl, block = soft_scene
    vid = tpl["default_version_id"]
    rid = tpl["regions"][0]["id"]
    response = client.post(
        f"/api/versions/{vid}/bindings",
        json={"region_id": rid, "block_id": block["id"], "line_break_mode": "soft"},
    )
    assert response.status_code == 200 and response.json()["line_break_mode"] == "soft"
    rendered = render_version(vid)
    doc = Document(BytesIO(rendered.data))
    assert len(doc.paragraphs) == 2
    assert doc.paragraphs[0].text == "前缀第一行\n第二行后缀"
    assert any(run.bold for run in doc.paragraphs[0].runs if run.text)
    assert doc.paragraphs[0]._p.xpath(".//w:br")
    assert list(iter_flow_paragraphs(rendered.data))[0][1] == doc.paragraphs[0].text
    assert rendered.items[0]["bbox"] is not None
    exported = client.post(f"/api/versions/{vid}/export", json={"confirm_large_overflow": True})
    assert exported.content == rendered.data
    copied = client.post(
        f"/api/templates/{tpl['id']}/versions", json={"name": "复制", "copy_from": vid}
    ).json()
    with get_conn() as conn:
        found_1 = bindings.get_binding(conn, copied["id"], rid)
        assert found_1 is not None
        assert found_1.line_break_mode == "soft"
    edit = client.post(
        f"/api/versions/{vid}/regions/{rid}/text",
        json={"content": "新第一行\n新第二行", "expected_content": block["content"]},
    )
    assert edit.status_code == 200
    with get_conn() as conn:
        found_2 = bindings.get_binding(conn, vid, rid)
        assert found_2 is not None
        assert found_2.line_break_mode == "soft"


def test_invalid_mode_does_not_create_binding(
    client: TestClient, soft_scene: tuple[dict, dict]
) -> None:
    tpl, block = soft_scene
    response = client.post(
        f"/api/versions/{tpl['default_version_id']}/bindings",
        json={
            "region_id": tpl["regions"][0]["id"],
            "block_id": block["id"],
            "line_break_mode": "unknown",
        },
    )
    assert response.status_code == 422
    with get_conn() as conn:
        assert bindings.list_bindings(conn, tpl["default_version_id"]) == []
