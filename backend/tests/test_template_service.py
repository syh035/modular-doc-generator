"""template_service 单测：校验链四分支、上传落库落盘、重复拦截、失败无残留。"""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from app.core.config import settings
from app.core.errors import AppError
from app.models.db import get_conn, init_db
from app.models.repositories import regions as regions_repo
from app.models.repositories import templates as templates_repo
from app.services.docx_parser import parse_placeholders  # noqa: F401  (保证导入链可用)
from app.services.template_service import upload_template

# OLE2 魔数（加密 DOCX / 旧 .doc 的文件头）
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


@pytest.fixture
def env(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """隔离环境：DB 与模板目录都指向 tmp_path。"""
    monkeypatch.setattr(settings, "db_path", tmp_path / "app.db")
    monkeypatch.setattr(settings, "templates_dir", tmp_path / "templates")
    with get_conn() as conn:
        init_db(conn)
    return tmp_path


def make_docx(label_paragraph: str = "{{姓名}}求职") -> bytes:
    doc = Document()
    doc.add_paragraph(label_paragraph)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_zip_without_document() -> bytes:
    """合法 zip 但缺 word/document.xml。"""
    buf = BytesIO()
    with ZipFile(buf, "w") as zf:
        zf.writestr("unrelated.txt", "占位")
    return buf.getvalue()


# ---- 校验链 ----


def test_reject_non_docx_extension(env: Path) -> None:
    for name in ("a.doc", "a.docm", "a.dotx", "无扩展名", ""):
        with pytest.raises(AppError) as ei:
            upload_template(name, make_docx())
        assert ei.value.code == "TEMPLATE_NOT_DOCX", name


def test_reject_ole_encrypted(env: Path) -> None:
    with pytest.raises(AppError) as ei:
        upload_template("加密.docx", OLE_MAGIC + b"\x00" * 64)
    assert ei.value.code == "TEMPLATE_ENCRYPTED"


def test_reject_corrupt_zip(env: Path) -> None:
    with pytest.raises(AppError) as ei:
        upload_template("损坏.docx", b"this is not a zip file")
    assert ei.value.code == "TEMPLATE_CORRUPT"


def test_reject_zip_missing_document_xml(env: Path) -> None:
    with pytest.raises(AppError) as ei:
        upload_template("缺正文.docx", make_zip_without_document())
    assert ei.value.code == "TEMPLATE_CORRUPT"


def test_reject_empty_file(env: Path) -> None:
    with pytest.raises(AppError) as ei:
        upload_template("空.docx", b"")
    assert ei.value.code == "TEMPLATE_CORRUPT"


# ---- 正常上传 ----


def test_upload_success(env: Path) -> None:
    data = make_docx("{{姓名}}的简历，邮箱{{邮箱}}")
    tpl, regions, reused = upload_template("我的模板.docx", data)

    assert tpl.filename == "我的模板.docx"
    assert tpl.status == "pending_review"
    assert tpl.storage_name == f"{tpl.id}_我的模板.docx"
    assert reused is False
    # 落盘校验
    stored = env / "templates" / tpl.storage_name
    assert stored.is_file() and stored.read_bytes() == data
    # 区域落库：决策 C——占位符 label 过词表推断类型（姓名→name / 邮箱→contact）
    assert [r.label for r in regions] == ["姓名", "邮箱"]
    assert [r.type for r in regions] == ["name", "contact"]
    assert [r.order_index for r in regions] == [0, 1]
    assert regions[0].placeholder == "{{姓名}}"
    # 状态机轨迹：建档 parsing 已被最终态覆盖（同步解析，终态即 pending_review）
    with get_conn() as conn:
        assert templates_repo.get_template(conn, tpl.id) is not None
        assert len(regions_repo.list_regions(conn, tpl.id)) == 2


def test_upload_sanitizes_path_in_filename(env: Path) -> None:
    """multipart 文件名带路径 → 只取纯文件名落盘。"""
    data = make_docx()
    tpl, _, _ = upload_template("../../etc/evil.docx", data)
    assert tpl.filename == "evil.docx"
    assert "/" not in tpl.storage_name
    assert (env / "templates" / tpl.storage_name).is_file()


def test_upload_placeholder_free_template(env: Path) -> None:
    """无占位符模板：解析成功 0 区域，仍进待校对（M5b 框选兜底）。"""
    tpl, regions, _ = upload_template("素模板.docx", make_docx("纯文本段落"))
    assert tpl.status == "pending_review"
    assert regions == []


# ---- 重复与同名 ----


def test_duplicate_content_reuses_existing(env: Path) -> None:
    """D10 指纹关联：同 sha256 重传 → 幂等返回已有模板，不新建记录。"""
    data = make_docx()
    tpl_first, _, reused_first = upload_template("第一份.docx", data)
    tpl_second, regions, reused = upload_template("换名同内容.docx", data)
    assert reused_first is False
    assert reused is True
    assert tpl_second.id == tpl_first.id
    assert tpl_second.filename == "第一份.docx"  # 沿用首传文件名
    assert len(regions) == 1  # 解析结果复用（make_docx 默认单占位符）
    with get_conn() as conn:
        assert len(templates_repo.list_templates(conn)) == 1


def test_image_only_template_rejected(env: Path) -> None:
    """M3b：正文无文本层（纯图片模板）→ TEMPLATE_IMAGE_ONLY，无残留。"""
    doc = Document()  # 默认模板仅空段落
    buf = BytesIO()
    doc.save(buf)
    with pytest.raises(AppError) as ei:
        upload_template("纯图片.docx", buf.getvalue())
    assert ei.value.code == "TEMPLATE_IMAGE_ONLY"
    assert ei.value.status_code == 400
    with get_conn() as conn:
        assert templates_repo.list_templates(conn) == []
    assert list((env / "templates").glob("*")) == []  # 落盘清残


def test_same_filename_different_content_allowed(env: Path) -> None:
    """PRD 4.2：同名模板允许重复导入，靠记录区分（id/时间戳）。"""
    tpl_a, _, _ = upload_template("简历模板.docx", make_docx("{{姓名}}A"))
    tpl_b, _, _ = upload_template("简历模板.docx", make_docx("{{姓名}}B"))
    assert tpl_a.id != tpl_b.id
    assert tpl_a.sha256 != tpl_b.sha256
    assert tpl_a.storage_name != tpl_b.storage_name
    assert tpl_a.created_at <= tpl_b.created_at  # 时间戳可区分先后
    with get_conn() as conn:
        assert len(templates_repo.list_templates(conn)) == 2


# ---- 失败无残留 ----


def test_failed_upload_leaves_no_trace(env: Path) -> None:
    for bad_name, bad_data in [
        ("坏扩展.txt", b"x"),
        ("加密.docx", OLE_MAGIC),
        ("损坏.docx", b"junk"),
    ]:
        with pytest.raises(AppError):
            upload_template(bad_name, bad_data)
    with get_conn() as conn:
        assert templates_repo.list_templates(conn) == []
    assert list((env / "templates").glob("*")) == []
