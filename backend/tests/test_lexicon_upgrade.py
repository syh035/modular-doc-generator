"""M15 历史候选升级：身份／绑定保留，只升级用户授权的待校对自动分类。"""

import json
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from app.core.config import settings
from app.models.db import get_conn
from app.models.repositories import bindings, blocks, regions
from app.services.template_service import upgrade_full_candidates, upgrade_pending_candidate_types


def upload(client: TestClient, texts: list[str]) -> tuple[dict, bytes]:
    doc = Document()
    for text in texts:
        doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    raw = buf.getvalue()
    response = client.post("/api/templates", files={"file": ("词表升级.docx", raw)})
    assert response.status_code == 201
    return response.json(), raw


def test_existing_pending_auto_candidates_upgrade_once_preserving_identity(
    client: TestClient,
) -> None:
    tpl, raw = upload(client, ["个人技能", "个人综述", "工作履历", "科研项目", "现居地"])
    ids = [r["id"] for r in tpl["regions"]]
    with get_conn() as conn:
        for rid in ids:
            regions.update_region(conn, rid, region_type="custom", confidence=0.4)
        regions.update_region(conn, ids[1], review_status="confirmed")
        regions.update_region(conn, ids[2], review_status="excluded")
        regions.update_region(conn, ids[3], bbox_source="manual", bbox={"page": 0, "y0": 1})
        block = blocks.create_block(conn, "现有块", "现有内容")
        bindings.upsert_binding(conn, tpl["default_version_id"], ids[0], block.id)
        conn.execute(
            "UPDATE template_candidate_scans SET revision = 2 WHERE template_id = ?", (tpl["id"],)
        )
    with get_conn() as conn:
        assert upgrade_pending_candidate_types(conn) == 2
        result = regions.list_regions(conn, tpl["id"])
        assert [r.id for r in result] == ids
        assert [r.type for r in result] == ["skills", "custom", "custom", "custom", "contact"]
        assert result[0].confidence == 0.9 and result[0].review_status == "pending"
        assert result[0].label == "个人技能"
        assert result[1].review_status == "confirmed"
        assert result[2].review_status == "excluded"
        assert result[3].bbox_json is not None
        assert result[3].bbox_source == "manual" and json.loads(result[3].bbox_json) == {
            "page": 0,
            "y0": 1,
        }
        found_1 = bindings.get_binding(conn, tpl["default_version_id"], ids[0])
        assert found_1 is not None
        assert found_1.block_id == block.id
        regions.update_region(conn, ids[0], region_type="custom")
    with get_conn() as conn:
        assert upgrade_pending_candidate_types(conn) == 0
        found_2 = regions.get_region(conn, ids[0])
        assert found_2 is not None
        assert found_2.type == "custom"
    assert (settings.templates_dir / tpl["storage_name"]).read_bytes() == raw


def test_same_paragraph_placeholders_are_reclassified_independently(client: TestClient) -> None:
    tpl, _ = upload(client, ["{{个人技能}} · {{个人综述}} · {{自定义}}"])
    ids = [r["id"] for r in tpl["regions"]]
    with get_conn() as conn:
        for rid in ids:
            regions.update_region(conn, rid, region_type="custom")
        conn.execute(
            "UPDATE template_candidate_scans SET revision = 2 WHERE template_id = ?", (tpl["id"],)
        )
    with get_conn() as conn:
        assert upgrade_pending_candidate_types(conn) == 2
        result = regions.list_regions(conn, tpl["id"])
        assert [r.type for r in result] == ["skills", "summary", "custom"]
        assert all(r.confidence == 1.0 for r in result)
        assert [r.placeholder for r in result] == ["{{个人技能}}", "{{个人综述}}", "{{自定义}}"]


def test_old_full_scan_still_backfills_newly_known_short_title(client: TestClient) -> None:
    tpl, _ = upload(client, ["个人技能", "一、教育背景", "工作经历"])
    with get_conn() as conn:
        conn.execute("DELETE FROM template_candidate_scans WHERE template_id = ?", (tpl["id"],))
        conn.execute("DELETE FROM regions WHERE template_id = ?", (tpl["id"],))
    with get_conn() as conn:
        # 旧词表标题曾被物理删除，不恢复；新增词表标题过去未识别，需补齐。
        assert upgrade_full_candidates(conn) == 2
        result = regions.list_regions(conn, tpl["id"])
        assert [r.label for r in result] == ["个人技能", "一、教育背景"]
        assert [r.type for r in result] == ["skills", "education"]
        assert upgrade_pending_candidate_types(conn) == 0
        assert upgrade_full_candidates(conn) == 0
