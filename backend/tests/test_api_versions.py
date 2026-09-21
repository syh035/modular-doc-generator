"""/api/versions 测试（M6a）：默认版本、绑定 CRUD、版本渲染（假 LO）+ 真实 LO 端到端。"""

from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.services.libreoffice import find_soffice


def make_docx(text: str = "电话：{{手机号}}，姓名：{{姓名}}") -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def upload(client: TestClient, data: bytes) -> dict:
    resp = client.post(
        "/api/templates",
        files={"file": ("模板.docx", data, "application/octet-stream")},
    )
    assert resp.status_code == 201
    return resp.json()


def setup_version(client: TestClient, data: bytes) -> tuple[int, list[dict]]:
    """上传模板 → 返回 (默认版本 id, 区域列表)。"""
    body = upload(client, data)
    assert body["default_version_id"] is not None
    return body["default_version_id"], body["regions"]


def make_block(client: TestClient, name: str, content: str) -> int:
    resp = client.post("/api/blocks", json={"name": name, "content": content})
    assert resp.status_code == 201
    return resp.json()["id"]


def bind(client: TestClient, version_id: int, region_id: int, block_id: int):
    return client.post(
        f"/api/versions/{version_id}/bindings",
        json={"region_id": region_id, "block_id": block_id},
    )


def fake_pdf_with(tmp_path: Path, text: str) -> Path:
    """造一份含指定文本的「渲染 PDF」（绕开 LO 验证替换-对齐链路）。"""
    pdf = tmp_path / f"fake_{abs(hash(text))}.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), text, fontname="china-s")
    doc.save(pdf)
    doc.close()
    return pdf


# ---- 默认版本 ----


def test_upload_creates_default_version(client: TestClient) -> None:
    body = upload(client, make_docx())
    assert body["default_version_id"] > 0
    detail = client.get(f"/api/templates/{body['id']}").json()
    assert detail["default_version_id"] == body["default_version_id"]


# ---- 绑定 CRUD ----


def test_bind_and_rebind_upsert(client: TestClient) -> None:
    vid, regions = setup_version(client, make_docx())
    rid = regions[0]["id"]
    b1 = make_block(client, "手机号A", "138")
    b2 = make_block(client, "手机号B", "139")

    resp = bind(client, vid, rid, b1)
    assert resp.status_code == 200
    assert resp.json()["block_id"] == b1
    assert resp.json()["status"] == "active"

    # 换绑：同端点覆盖
    resp = bind(client, vid, rid, b2)
    assert resp.status_code == 200
    assert resp.json()["block_id"] == b2

    items = client.get(f"/api/versions/{vid}/bindings").json()["bindings"]
    assert len(items) == 1  # UNIQUE(version, region)：仍是 1 条
    assert items[0]["block_id"] == b2
    assert items[0]["block_name"] == "手机号B"


def test_unbind_204_then_404(client: TestClient) -> None:
    vid, regions = setup_version(client, make_docx())
    rid = regions[0]["id"]
    bid = make_block(client, "手机号", "138")
    assert bind(client, vid, rid, bid).status_code == 200

    resp = client.delete(f"/api/versions/{vid}/bindings/{rid}")
    assert resp.status_code == 204
    resp = client.delete(f"/api/versions/{vid}/bindings/{rid}")
    assert resp.status_code == 404
    assert resp.json()["error"]["code"] == "BINDING_NOT_FOUND"


def test_bind_validations(client: TestClient) -> None:
    vid, regions = setup_version(client, make_docx())
    rid = regions[0]["id"]

    # 版本不存在
    assert bind(client, 999, rid, 1).status_code == 404
    # 区域不存在
    resp = bind(client, vid, 9999, 1)
    assert resp.status_code == 404
    assert resp.json()["error"]["code"] == "REGION_NOT_FOUND"
    # 块不存在
    resp = bind(client, vid, rid, 9999)
    assert resp.status_code == 404
    assert resp.json()["error"]["code"] == "BLOCK_NOT_FOUND"

    # 区域不属于该版本所属模板：上传第二个模板做交叉
    other = upload(client, make_docx("别的{{字段}}"))
    other_vid = other["default_version_id"]
    bid = make_block(client, "交叉块", "内容")
    resp = bind(client, other_vid, rid, bid)  # 模板1的区域绑到模板2的版本
    assert resp.status_code == 400
    assert resp.json()["error"]["code"] == "REGION_TEMPLATE_MISMATCH"


# ---- 版本渲染（假 LO：验证替换 + 对齐 + 绑定态编排）----


def _patch_fake_convert(
    monkeypatch: pytest.MonkeyPatch, pdf_by_text: dict[str, Path]
) -> list[bytes]:
    """按成品 DOCX 正文文本选假 PDF（DOCX 是 zip，须解析后匹配）；记录被转换字节。"""
    seen: list[bytes] = []

    def convert(path: Path) -> Path:
        data = path.read_bytes()
        seen.append(data)
        text = "\n".join(p.text for p in Document(BytesIO(data)).paragraphs)
        for needle, pdf in pdf_by_text.items():
            if needle in text:
                return pdf
        raise AssertionError(f"无匹配假 PDF：{path}")

    monkeypatch.setattr("app.services.libreoffice._manager", SimpleNamespace(convert=convert))
    return seen


def test_version_preview_and_overlay(
    client: TestClient, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """绑定 → 版本预览 PDF = 替换后成品；overlay 带绑定态与替换后 bbox。"""
    tpl_data = make_docx("电话：{{手机号}}，姓名：{{姓名}}")
    vid, regions = setup_version(client, tpl_data)
    by_label = {r["label"]: r for r in regions}
    b_phone = make_block(client, "手机号", "13800138000")
    assert bind(client, vid, by_label["手机号"]["id"], b_phone).status_code == 200
    # 姓名不绑 → 保留占位符

    # 假 PDF 需含替换后文本（几何对齐源）：绑定行 + 保留的占位符行；
    # 模板本体转换（M7 溢出测量触发 ensure 补 bbox）用原文假 PDF
    fake = fake_pdf_with(
        tmp_path, "电话：13800138000，姓名：{{姓名}}"
    )
    fake_tpl = fake_pdf_with(tmp_path, "电话：{{手机号}}，姓名：{{姓名}}")
    _patch_fake_convert(
        monkeypatch, {"电话：13800138000": fake, "电话：{{手机号}}": fake_tpl}
    )

    resp = client.get(f"/api/versions/{vid}/preview")
    assert resp.status_code == 200
    assert resp.content.startswith(b"%PDF")

    overlay = client.get(f"/api/versions/{vid}/overlay").json()
    items = {r["label"]: r for r in overlay["regions"]}
    phone = items["手机号"]
    assert phone["binding"]["block_id"] == b_phone
    assert phone["binding"]["block_name"] == "手机号"
    assert phone["binding"]["status"] == "active"
    assert phone["bbox"] is not None  # 替换后文本对齐成功
    assert items["姓名"]["binding"] is None
    assert items["姓名"]["bbox"] is not None  # 未绑定区域也原位对齐


def test_version_render_unbound_keeps_placeholder(
    client: TestClient, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """零绑定时版本渲染 = 原文（占位符保留），仍是合法 PDF。"""
    vid, _ = setup_version(client, make_docx())
    fake = fake_pdf_with(tmp_path, "电话：{{手机号}}，姓名：{{姓名}}")
    seen = _patch_fake_convert(monkeypatch, {"{{手机号}}": fake})

    assert client.get(f"/api/versions/{vid}/preview").status_code == 200
    doc = Document(BytesIO(seen[0]))  # 送入 LO 的成品 = 原文
    assert doc.paragraphs[0].text == "电话：{{手机号}}，姓名：{{姓名}}"


def test_version_404(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "app.services.libreoffice._manager", SimpleNamespace(convert=lambda p: p)
    )
    for path in ("/api/versions/999/preview", "/api/versions/999/overlay"):
        resp = client.get(path)
        assert resp.status_code == 404
        assert resp.json()["error"]["code"] == "VERSION_NOT_FOUND"


# ---- 真实 LibreOffice 端到端 ----


@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice 未安装")
def test_version_render_real_lo(client: TestClient) -> None:
    """里程碑 1 风险链核心：绑定块 → 替换 → LO 渲染 → 文本与坐标可提取。

    python-docx 默认模板 run 无 eastAsia（M4 事实）→ 显式设置，
    保证 CJK 字形渲染（验证墨迹须真实 LO，此处以文本提取 + bbox 为准）。
    """
    doc = Document()
    p = doc.add_paragraph("电话：{{手机号}}")
    run = p.runs[0]
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "宋体"
    )
    buf = BytesIO()
    doc.save(buf)
    tpl_data = buf.getvalue()

    vid, regions = setup_version(client, tpl_data)
    bid = make_block(client, "手机号块", "13800138000")
    assert bind(client, vid, regions[0]["id"], bid).status_code == 200

    resp = client.get(f"/api/versions/{vid}/preview")
    assert resp.status_code == 200
    assert resp.content.startswith(b"%PDF")

    # 成品 PDF 必须含替换文本、不含占位符（文本层验证；墨迹验证见手动验收）
    with pymupdf.open(stream=resp.content) as pdf:
        text = pdf[0].get_text()
    assert "13800138000" in text
    assert "{{手机号}}" not in text

    overlay = client.get(f"/api/versions/{vid}/overlay").json()
    region = overlay["regions"][0]
    assert region["bbox"] is not None
    assert region["binding"]["status"] == "active"
