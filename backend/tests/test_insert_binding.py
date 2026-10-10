"""Insertion leaves original anchors intact and preserves layout settings."""

from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from fastapi.testclient import TestClient

from app.core.config import settings
from app.models.db import get_conn
from app.models.repositories import regions as regions_repo
from app.services.replacement import apply_replacements


def scene(client: TestClient, suffix: str = "") -> dict:
    doc = Document()
    p = doc.add_paragraph("原文{{姓名}}" + suffix)
    p.runs[0].bold = True
    ppr = p._p.get_or_add_pPr()
    num = OxmlElement("w:numPr")
    ppr.append(num)
    doc.add_paragraph("{{电话}}")
    buf = BytesIO()
    doc.save(buf)
    response = client.post("/api/templates", files={"file": ("插入.docx", buf.getvalue())})
    assert response.status_code == 201
    return response.json()


def test_insert_soft_before_after_blank_preserves_identity(client: TestClient) -> None:
    tpl = scene(client)
    with get_conn() as conn:
        regions = regions_repo.list_regions(conn, tpl["id"])
    first = next(r for r in regions if r.placeholder == "{{姓名}}")
    last = next(r for r in regions if r.placeholder == "{{电话}}")
    source = (settings.templates_dir / tpl["storage_name"]).read_bytes()
    out = apply_replacements(
        source,
        regions,
        {first.id: "前一行\n前二行", last.id: ""},
        positions={first.id: "before", last.id: "after"},
        line_break_modes={first.id: "soft"},
        blank_region_ids={last.id},
    )
    doc = Document(BytesIO(out.data))
    assert [p.text for p in doc.paragraphs] == ["前一行\n前二行", "原文{{姓名}}", "{{电话}}", ""]
    assert doc.paragraphs[0].runs[0].bold
    assert not doc.paragraphs[0]._p.xpath("./w:pPr/w:numPr")
    assert doc.paragraphs[1]._p.xpath("./w:pPr/w:numPr")
    assert out.region_paths[first.id] == [1]
    assert out.region_paths[last.id] == [2]
    assert out.region_clone_paths == {first.id: [[0]], last.id: [[3]]}
    assert (settings.templates_dir / tpl["storage_name"]).read_bytes() == source


def test_insert_copy_and_migration_keep_mode_position(client: TestClient) -> None:
    from app.services.render_service import render_version

    source = scene(client)
    target = scene(client, "新模板")
    vid = source["default_version_id"]
    rid = next(r["id"] for r in source["regions"] if r["placeholder"] == "{{姓名}}")
    block = client.post(
        "/api/blocks", json={"name": "插入内容", "content": "插入甲\n插入乙"}
    ).json()
    bound = client.post(
        f"/api/versions/{vid}/bindings",
        json={
            "region_id": rid,
            "block_id": block["id"],
            "line_break_mode": "soft",
            "position": "after",
        },
    )
    assert bound.status_code == 200 and bound.json()["position"] == "after"
    rendered = render_version(vid)
    assert [p.text for p in Document(BytesIO(rendered.data)).paragraphs] == [
        "原文{{姓名}}",
        "插入甲\n插入乙",
        "{{电话}}",
    ]
    assert all(r["bbox"] is not None for r in rendered.items)
    copied = client.post(
        f"/api/templates/{source['id']}/versions", json={"name": "副本", "copy_from": vid}
    ).json()
    binding = client.get(f"/api/versions/{copied['id']}/bindings").json()["bindings"][0]
    assert binding["position"] == "after" and binding["line_break_mode"] == "soft"
    plan = client.post(
        f"/api/templates/{target['id']}/migrate/plan", json={"source_version_id": vid}
    ).json()
    row = plan["auto"][0]
    assert row["position"] == "after" and row["line_break_mode"] == "soft"
    response = client.post(
        f"/api/templates/{target['id']}/migrate/apply",
        json={
            "source_version_id": vid,
            "bindings": [
                {
                    "region_id": row["target_region_id"],
                    "block_id": block["id"],
                    "position": row["position"],
                    "line_break_mode": row["line_break_mode"],
                }
            ],
        },
    )
    assert response.status_code == 201
    binding = client.get(f"/api/versions/{target['default_version_id']}/bindings").json()[
        "bindings"
    ][0]
    assert binding["position"] == "after" and binding["line_break_mode"] == "soft"
    export = client.post(f"/api/versions/{vid}/export", json={"confirm_large_overflow": True})
    assert export.content == rendered.data
