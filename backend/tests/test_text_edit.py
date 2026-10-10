"""版本文字微调：隔离、显式共享、并发保护、事务与真实渲染。"""

from io import BytesIO

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.core.config import settings
from app.models.db import get_conn
from app.models.repositories import bindings, blocks, regions, versions
from app.services import text_edit_service
from app.services.libreoffice import find_soffice
from app.services.render_service import render_version


@pytest.fixture
def scene(client: TestClient) -> tuple[TestClient, dict, bytes]:
    doc = Document()
    doc.add_paragraph().add_run("原文字").bold = True
    buf = BytesIO()
    doc.save(buf)
    raw = buf.getvalue()
    tpl = client.post("/api/templates", files={"file": ("文字.docx", raw)}).json()
    return client, tpl, raw


def save(
    scene: tuple[TestClient, dict, bytes],
    text: str = "微调文字",
    sync: bool = False,
    expected: str | None = "原文字",
    vid: int | None = None,
) -> httpx.Response:
    c, tpl, _ = scene
    return c.post(
        f"/api/versions/{vid or tpl['default_version_id']}/regions/{tpl['regions'][0]['id']}/text",
        json={"content": text, "sync_block": sync, "expected_content": expected},
    )


def test_unbound_creates_block_without_touching_original(
    scene: tuple[TestClient, dict, bytes],
) -> None:
    c, tpl, raw = scene
    response = save(scene)
    assert response.status_code == 200
    with get_conn() as conn:
        block = blocks.get_block(conn, response.json()["block_id"])
        assert block is not None
        assert block.content == "微调文字"
        found_1 = bindings.get_binding(conn, tpl["default_version_id"], tpl["regions"][0]["id"])
        assert found_1 is not None
        assert found_1.block_id == block.id
    assert (settings.templates_dir / tpl["storage_name"]).read_bytes() == raw


@pytest.mark.parametrize("sync", [False, True])
def test_shared_block_only_changes_when_explicitly_synced(
    scene: tuple[TestClient, dict, bytes], sync: bool
) -> None:
    _, tpl, _ = scene
    vid = tpl["default_version_id"]
    rid = tpl["regions"][0]["id"]
    original_id = save(scene).json()["block_id"]
    with get_conn() as conn:
        other = versions.create_version(conn, tpl["id"], "另一版本")
        bindings.upsert_binding(conn, other.id, rid, original_id)
        conn.execute("INSERT INTO tags (name, created_at) VALUES ('标签', 'now')")
        tag_id = conn.execute("SELECT id FROM tags").fetchone()[0]
        conn.execute(
            "INSERT INTO block_tags (block_id, tag_id) VALUES (?, ?)", (original_id, tag_id)
        )
    response = save(scene, "再次微调", sync, "微调文字")
    assert response.status_code == 200
    new_id = response.json()["block_id"]
    assert (new_id == original_id) is sync
    with get_conn() as conn:
        found_2 = blocks.get_block(conn, original_id)
        assert found_2 is not None
        assert found_2.content == ("再次微调" if sync else "微调文字")
        found_3 = blocks.get_block(conn, new_id)
        assert found_3 is not None
        assert found_3.content == "再次微调"
        found_4 = bindings.get_binding(conn, other.id, rid)
        assert found_4 is not None
        assert found_4.block_id == original_id
        found_5 = bindings.get_binding(conn, vid, rid)
        assert found_5 is not None
        assert found_5.block_id == new_id
        assert (
            conn.execute("SELECT tag_id FROM block_tags WHERE block_id = ?", (new_id,)).fetchone()[
                0
            ]
            == tag_id
        )


@pytest.mark.parametrize(
    "content,expected,status,code",
    [
        ("", "原文字", 400, "BLOCK_INVALID"),
        (" " * 10, "原文字", 400, "BLOCK_INVALID"),
        ("字" * 5001, "原文字", 400, "BLOCK_INVALID"),
        ("新内容", "过期内容", 409, "REGION_TEXT_CONFLICT"),
    ],
)
def test_invalid_or_stale_edit_leaves_no_blocks(
    scene: tuple[TestClient, dict, bytes],
    content: str,
    expected: str | None,
    status: int,
    code: str,
) -> None:
    response = save(scene, content, expected=expected)
    assert response.status_code == status
    assert response.json()["error"]["code"] == code
    with get_conn() as conn:
        assert conn.execute("SELECT COUNT(*) FROM blocks").fetchone()[0] == 0


def test_excluded_and_wrong_template_refused(scene: tuple[TestClient, dict, bytes]) -> None:
    c, tpl, _ = scene
    rid = tpl["regions"][0]["id"]
    with get_conn() as conn:
        regions.update_region(conn, rid, review_status="excluded")
    assert save(scene).json()["error"]["code"] == "REGION_EXCLUDED"
    assert save(scene, vid=9999).status_code == 404
    assert save(scene, sync=True).status_code == 400


def test_sync_requires_live_binding(scene: tuple[TestClient, dict, bytes]) -> None:
    assert save(scene, sync=True).json()["error"]["code"] == "BLOCK_NOT_FOUND"


def test_binding_failure_rolls_back_block(
    scene: tuple[TestClient, dict, bytes], monkeypatch: pytest.MonkeyPatch
) -> None:
    _, tpl, _ = scene

    def fail(*args: object) -> None:
        raise RuntimeError("injected binding failure")

    monkeypatch.setattr(bindings, "upsert_binding", fail)
    with pytest.raises(RuntimeError):
        text_edit_service.save_text(
            tpl["default_version_id"], tpl["regions"][0]["id"], "新内容", False
        )
    with get_conn() as conn:
        assert conn.execute("SELECT COUNT(*) FROM blocks").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM bindings").fetchone()[0] == 0


@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice 未安装")
def test_edit_preview_and_export_share_text_and_style(
    scene: tuple[TestClient, dict, bytes],
) -> None:
    c, tpl, raw = scene
    vid = tpl["default_version_id"]
    rendered_before = render_version(vid)
    assert rendered_before.items[0]["current_text"] == "原文字"
    assert save(scene).status_code == 200
    rendered = render_version(vid)
    assert rendered.items[0]["current_text"] == "微调文字"
    p = Document(BytesIO(rendered.data)).paragraphs[0]
    assert p.text == "微调文字" and any(run.bold for run in p.runs if run.text)
    exported = c.post(f"/api/versions/{vid}/export", json={"confirm_large_overflow": True})
    assert exported.status_code == 200 and exported.content == rendered.data
    assert (settings.templates_dir / tpl["storage_name"]).read_bytes() == raw
