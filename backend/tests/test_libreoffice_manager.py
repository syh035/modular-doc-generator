"""LibreOfficeManager 单测：假 soffice 驱动缓存/超时回收/失败路径 + 真实转换集成。

假 soffice 是一个可执行 python 脚本：解析 --outdir 与末位输入路径，
按输入 stem 落一个假 PDF；FAKE_HANG=1 挂起 300s（超时回收用），
FAKE_FAIL=1 退出码 3（失败路径用），成功时向输入旁 fake_calls 计数。
"""

import stat
from io import BytesIO
from pathlib import Path

import pymupdf
import pytest
from docx import Document

from app.core.errors import RENDER_FAILED, RENDER_TIMEOUT, AppError
from app.services.libreoffice import LibreOfficeManager, find_soffice

_FAKE_SOFFICE = '''#!/usr/bin/env python3
import os
import sys
import time
from pathlib import Path

args = sys.argv[1:]
outdir = Path(args[args.index("--outdir") + 1])
inp = Path(args[-1])
marker = inp.parent / "fake_calls"
if os.environ.get("FAKE_HANG"):
    time.sleep(300)
if os.environ.get("FAKE_FAIL"):
    sys.exit(3)
(outdir / (inp.stem + ".pdf")).write_bytes(b"%PDF-fake\\n%M4")
with marker.open("a") as f:
    f.write("x")
'''


@pytest.fixture
def fake_soffice(tmp_path: Path) -> Path:
    p = tmp_path / "fake_soffice"
    p.write_text(_FAKE_SOFFICE)
    p.chmod(p.stat().st_mode | stat.S_IEXEC)
    return p


def make_manager(fake: Path, tmp_path: Path, timeout: float = 30.0) -> LibreOfficeManager:
    return LibreOfficeManager(
        str(fake),
        profile_dir=tmp_path / "prof",
        cache_dir=tmp_path / "cache",
        timeout_seconds=timeout,
    )


def make_docx_file(tmp_path: Path, name: str = "模板.docx") -> Path:
    doc = Document()
    doc.add_paragraph("张三的简历")
    buf = BytesIO()
    doc.save(buf)
    p = tmp_path / name
    p.write_bytes(buf.getvalue())
    return p


# ---- 假 soffice：成功 / 缓存 / 失败 / 超时回收 ----


def test_convert_ok_and_cached(fake_soffice: Path, tmp_path: Path) -> None:
    m = make_manager(fake_soffice, tmp_path)
    docx = make_docx_file(tmp_path)
    pdf1 = m.convert(docx)
    assert pdf1.is_file()
    assert pdf1.read_bytes().startswith(b"%PDF-fake")
    # 同内容二连：命中缓存，soffice 只跑一次
    pdf2 = m.convert(docx)
    assert pdf2 == pdf1
    assert m.conversions == 1
    assert (tmp_path / "fake_calls").read_text() == "x"
    assert not list((tmp_path / "cache" / "tmp").iterdir())  # 工作目录已清


def test_convert_failure_raises_render_failed(
    fake_soffice: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("FAKE_FAIL", "1")
    m = make_manager(fake_soffice, tmp_path)
    with pytest.raises(AppError) as ei:
        m.convert(make_docx_file(tmp_path))
    assert ei.value.code == RENDER_FAILED
    assert not list((tmp_path / "cache").glob("*.pdf"))  # 失败不留缓存残迹


def test_convert_timeout_kills_and_recovers(
    fake_soffice: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("FAKE_HANG", "1")
    m = make_manager(fake_soffice, tmp_path, timeout=0.5)
    with pytest.raises(AppError) as ei:
        m.convert(make_docx_file(tmp_path))
    assert ei.value.code == RENDER_TIMEOUT
    assert not (tmp_path / "prof").exists()  # profile 已清理（防残留锁）
    # 回收后可恢复：同 manager 下一次转换正常
    monkeypatch.delenv("FAKE_HANG")
    assert m.convert(make_docx_file(tmp_path)).is_file()


def test_different_content_no_cache_hit(fake_soffice: Path, tmp_path: Path) -> None:
    m = make_manager(fake_soffice, tmp_path)
    m.convert(make_docx_file(tmp_path, "a.docx"))
    doc2 = make_docx_file(tmp_path, "b.docx")
    doc2.write_bytes(doc2.read_bytes() + b"changed")  # 内容不同 → sha 不同
    m.convert(doc2)
    assert m.conversions == 2
    assert len(list((tmp_path / "cache").glob("*.pdf"))) == 2


# ---- 真实 LibreOffice 集成（未安装则跳过）----


@pytest.fixture(scope="module")
def _real_profile(tmp_path_factory: pytest.TempPathFactory) -> Path:
    return tmp_path_factory.mktemp("lo_profile")


@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice 未安装")
def test_real_convert_and_cache(tmp_path: Path, _real_profile: Path) -> None:
    soffice = find_soffice()
    assert soffice is not None
    m = LibreOfficeManager(
        soffice,
        profile_dir=_real_profile,
        cache_dir=tmp_path / "cache",
        timeout_seconds=60.0,  # 首次转换含 profile 冷启动，放宽
        fontconfig_file=tmp_path / "fc" / "fonts.conf",
    )
    doc = Document()
    doc.add_paragraph("张三的简历")
    doc.add_paragraph("电话：{{手机号}}")
    docx = tmp_path / "真模板.docx"
    buf = BytesIO()
    doc.save(buf)
    docx.write_bytes(buf.getvalue())

    pdf = m.convert(docx)
    assert pdf.read_bytes().startswith(b"%PDF")  # 真 PDF 魔数
    assert pdf.suffix == ".pdf"
    # P17 回归守卫：无字体设置的文档中文也必须有墨迹（fontconfig 修复生效）
    with pymupdf.open(pdf) as d:
        pix = d[0].get_pixmap(dpi=150)
        assert any(b < 200 for b in pix.samples), "页面无墨迹：中文渲染空白的 P17 复发"
    # 二连命中缓存
    again = m.convert(docx)
    assert again == pdf
    assert m.conversions == 1
