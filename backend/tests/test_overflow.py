"""溢出测量测试（M7）：纯函数 + 固定行高构造 + overlay API（真实 LO 端到端）。"""

from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

from app.services.docx_parser import iter_flow_paragraph_elements
from app.services.libreoffice import find_soffice
from app.services.overflow import (
    detect_clipped,
    detect_fixed_rows,
    measure_overflow,
    overflow_from_json,
)
from app.services.pdf_geometry import PdfLine

_EAST = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia"


# ---- measure_overflow（D4 分级）----


def _new_bbox(h: float) -> tuple[float, float, float, float]:
    return (10.0, 100.0, 200.0, 100.0 + h)


_ORIG = {"y0": 100.0, "y1": 120.0}  # 原高 20


def test_measure_no_overflow_returns_none() -> None:
    assert measure_overflow(_ORIG, _new_bbox(15.0), 0.5) is None  # 缩短
    assert measure_overflow(_ORIG, _new_bbox(20.0), 0.5) is None  # 持平


def test_measure_small_and_boundary() -> None:
    info = measure_overflow(_ORIG, _new_bbox(25.0), 0.5)
    assert info is not None
    assert info.level == "small"
    # D4 边界：ratio 恰等阈值属小超出（仅 > threshold 为大超出）
    info = measure_overflow(_ORIG, _new_bbox(30.0), 0.5)
    assert info is not None
    assert info.level == "small"
    assert info.ratio == 0.5


def test_measure_large_beyond_threshold() -> None:
    info = measure_overflow(_ORIG, _new_bbox(31.0), 0.5)
    assert info is not None
    assert info.level == "large"
    assert info.clipped is False
    assert info.orig_height == 20.0
    assert info.new_height == 31.0


def test_measure_clipped_forces_large() -> None:
    """P6：裁剪时 bbox 不长高（ratio≈0），仍按 large 级警示。"""
    info = measure_overflow(_ORIG, _new_bbox(20.0), 0.5, clipped=True, fixed_row=True)
    assert info is not None
    assert info.level == "large"
    assert info.clipped is True
    assert info.fixed_row is True


def test_measure_unmeasurable() -> None:
    assert measure_overflow(None, _new_bbox(30.0), 0.5) is None
    assert measure_overflow(_ORIG, None, 0.5) is None
    assert measure_overflow({"y0": 120.0, "y1": 120.0}, _new_bbox(30.0), 0.5) is None
    assert measure_overflow({}, _new_bbox(30.0), 0.5) is None


def test_overflow_from_json() -> None:
    assert overflow_from_json(None) is None
    assert overflow_from_json("") is None
    assert overflow_from_json("not json") is None
    assert overflow_from_json('{"y0": 1.0}') == {"y0": 1.0}
    assert overflow_from_json("[1, 2]") is None  # 非 dict


# ---- detect_fixed_rows（P6：hRule="exact" 才算固定行高）----


def _set_row_height(row: object, val: str, rule: str) -> None:
    tr_pr = row._tr.get_or_add_trPr()  # type: ignore[attr-defined]
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), val)
    h.set(qn("w:hRule"), rule)
    tr_pr.append(h)


def _table_docx_bytes(rule0: str | None, rule1: str | None) -> bytes:
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    if rule0:
        _set_row_height(table.rows[0], "400", rule0)
    if rule1:
        _set_row_height(table.rows[1], "400", rule1)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _cell_para_paths(data: bytes) -> list[list[int]]:
    doc = Document(BytesIO(data))
    paths = []
    for anchor, _el in iter_flow_paragraph_elements(doc.element.body):
        if anchor.get("kind") == "cell_p":
            path = anchor["path"]
            assert isinstance(path, list)
            paths.append(path)
    assert len(paths) == 2  # 两行单元格各一个（空）段落
    return paths


def test_detect_fixed_rows_exact_only() -> None:
    data = _table_docx_bytes("exact", "atLeast")
    p0, p1 = _cell_para_paths(data)
    assert detect_fixed_rows(data, {1: p0, 2: p1}) == {1}


def test_detect_fixed_rows_default_not_fixed() -> None:
    data = _table_docx_bytes(None, None)
    p0, p1 = _cell_para_paths(data)
    assert detect_fixed_rows(data, {1: p0, 2: p1}) == set()


def test_detect_fixed_rows_empty_paths() -> None:
    assert detect_fixed_rows(b"whatever", {}) == set()


# ---- detect_clipped（文本完整性比对）----


def _line(text: str) -> PdfLine:
    return PdfLine(0, text, (0.0, 0.0, 100.0, 12.0))


def test_detect_clipped_missing_line() -> None:
    lines = [_line("负责核心模块"), _line("性能优化落地")]
    assert detect_clipped(lines, {1: "负责核心模块\n性能优化落地"}, {1}) == set()
    assert detect_clipped(lines, {1: "负责核心模块\n第三行被裁剪"}, {1}) == {1}


def test_detect_clipped_wrapped_line_still_found() -> None:
    """内容行在 PDF 中折行：空白剔除后拼接仍是连续子串。"""
    lines = [_line("负责核心模块开发与性能"), _line("优化落地")]
    assert detect_clipped(lines, {1: "负责核心模块开发与性能优化落地"}, {1}) == set()


def test_detect_clipped_skips_unbound_and_empty() -> None:
    lines = [_line("第一行")]
    assert detect_clipped(lines, {1: "缺失内容\n第二行"}, set()) == set()
    assert detect_clipped(lines, {}, {1}) == set()


# ---- overlay API 端到端（真实 LO）----


def upload(client: TestClient, data: bytes) -> dict:
    resp = client.post(
        "/api/templates",
        files={"file": ("模板.docx", data, "application/octet-stream")},
    )
    assert resp.status_code == 201
    return resp.json()


def setup_version(client: TestClient, data: bytes) -> tuple[int, list[dict]]:
    body = upload(client, data)
    assert body["default_version_id"] is not None
    return body["default_version_id"], body["regions"]


def make_block(client: TestClient, name: str, content: str) -> int:
    resp = client.post("/api/blocks", json={"name": name, "content": content})
    assert resp.status_code == 201
    return resp.json()["id"]


def bind(client: TestClient, version_id: int, region_id: int, block_id: int):
    return client.post(
        f"/api/versions/{version_id}/bindings",
        json={"region_id": region_id, "block_id": block_id},
    )


def _cjk_run(doc: Document, text: str):
    p = doc.add_paragraph(text)
    run = p.runs[0]
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(_EAST, "宋体")
    return p


@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice 未安装")
def test_overlay_overflow_large_reflow(client: TestClient) -> None:
    """普通段落绑定多行内容 → LO 顺延重排，高度对比分级 large；未绑定区域无报告。"""
    doc = Document()
    _cjk_run(doc, "项目经历：{{项目经历}}")
    _cjk_run(doc, "姓名：{{姓名}}")
    buf = BytesIO()
    doc.save(buf)

    vid, regions = setup_version(client, buf.getvalue())
    assert len(regions) == 2
    rid = regions[0]["id"]
    bid = make_block(
        client,
        "项目块",
        "独立负责数据迁移平台\n完成核心链路重构\n性能提升三倍\n沉淀运维手册\n输出季度复盘",
    )
    assert bind(client, vid, rid, bid).status_code == 200

    overlay = client.get(f"/api/versions/{vid}/overlay").json()
    items = {r["label"]: r for r in overlay["regions"]}
    ov = items["项目经历"]["overflow"]
    assert ov is not None
    assert ov["level"] == "large"
    assert ov["ratio"] > 0.5
    assert ov["clipped"] is False
    assert ov["fixed_row"] is False
    assert ov["new_height"] > ov["orig_height"]
    assert items["姓名"]["overflow"] is None  # 未绑定不测量


@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice 未安装")
def test_overlay_overflow_clipped_in_fixed_row(client: TestClient) -> None:
    """P6：固定行高（exact 12pt）单元格绑定多行内容 → 被裁剪，clipped=true。"""
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell_p = table.cell(0, 0).paragraphs[0]
    run = cell_p.add_run("{{自我评价}}")
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(_EAST, "宋体")
    _set_row_height(table.rows[0], "240", "exact")
    buf = BytesIO()
    doc.save(buf)

    vid, regions = setup_version(client, buf.getvalue())
    bid = make_block(client, "评价块", "第一行评价内容\n第二行评价内容\n第三行评价内容")
    assert bind(client, vid, regions[0]["id"], bid).status_code == 200

    overlay = client.get(f"/api/versions/{vid}/overlay").json()
    ov = overlay["regions"][0]["overflow"]
    assert ov is not None
    assert ov["fixed_row"] is True
    assert ov["clipped"] is True
    assert ov["level"] == "large"
