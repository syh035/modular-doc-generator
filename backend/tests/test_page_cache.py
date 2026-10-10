"""Page equality includes resources/layout; cached results cannot be mutated by clients."""

from pathlib import Path

import pymupdf

from app.services.page_cache import page_fingerprints


def make_pdf(path: Path, first: str, *, rotate: bool = False, color: bool = False) -> None:
    with pymupdf.open() as doc:
        page = doc.new_page()
        page.insert_text((72, 72), first, color=(1, 0, 0) if color else (0, 0, 0))
        if rotate:
            page.set_rotation(90)
        doc.new_page().insert_text((72, 72), "Stable second page")
        doc.save(path)


def test_only_changed_page_differs_and_metadata_is_ignored(tmp_path: Path) -> None:
    a, b, c = [tmp_path / name for name in ("a.pdf", "b.pdf", "c.pdf")]
    make_pdf(a, "Old content")
    make_pdf(b, "New content")
    make_pdf(c, "Old content")
    sha_a, hashes_a = page_fingerprints(a)
    sha_b, hashes_b = page_fingerprints(b)
    assert sha_a != sha_b and hashes_a[0] != hashes_b[0]
    assert hashes_a[1] == hashes_b[1]
    assert page_fingerprints(c)[1] == hashes_a
    hashes_a.clear()
    assert len(page_fingerprints(a)[1]) == 2


def test_rotation_and_color_invalidate_page(tmp_path: Path) -> None:
    paths = [tmp_path / f"{i}.pdf" for i in range(3)]
    for path, rotate, color in zip(paths, [False, True, False], [False, False, True], strict=True):
        make_pdf(path, "Same text", rotate=rotate, color=color)
    hashes = [page_fingerprints(path)[1] for path in paths]
    assert len({h[0] for h in hashes}) == 3
    assert len({h[1] for h in hashes}) == 1


def test_same_path_new_pdf_does_not_reuse_old_hashes(tmp_path: Path) -> None:
    path = tmp_path / "preview.pdf"
    make_pdf(path, "Old content")
    old = page_fingerprints(path)
    path.unlink()
    make_pdf(path, "New content")
    assert page_fingerprints(path)[1][0] != old[1][0]


def test_unused_shared_font_glyphs_do_not_invalidate_visible_page(tmp_path: Path) -> None:
    # Actual LO subset order changes while page 2's visible glyph outlines stay equal.
    from io import BytesIO

    from docx import Document

    from app.services.libreoffice import LibreOfficeManager, find_soffice

    soffice = find_soffice()
    if soffice is None:
        import pytest

        pytest.skip("LibreOffice 未安装")
    # Never share the running application profile with a test subprocess.
    manager = LibreOfficeManager(
        soffice,
        profile_dir=tmp_path / "profile",
        cache_dir=tmp_path / "cache",
        timeout_seconds=120,
        fontconfig_file=tmp_path / "fontconfig/fonts.conf",
    )
    fingerprints = []
    for content in ["ABC", "CBA"]:
        doc = Document()
        doc.add_paragraph(content)
        doc.add_page_break()
        doc.add_paragraph("Second page remains stable")
        buf = BytesIO()
        doc.save(buf)
        path = tmp_path / f"{content}.docx"
        path.write_bytes(buf.getvalue())
        fingerprints.append(page_fingerprints(manager.convert(path))[1])
    assert fingerprints[0][0] != fingerprints[1][0]
    assert fingerprints[0][1] == fingerprints[1][1]
