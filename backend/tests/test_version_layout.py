"""Version paragraph removal/blank rendering must preserve anchors and originals."""

from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.core.config import settings
from app.models.db import get_conn
from app.models.repositories import regions as regions_repo
from app.services.replacement import apply_replacements


@pytest.fixture
def layout_scene(client: TestClient) -> dict:
    doc = Document()
    doc.add_paragraph("前文")
    p = doc.add_paragraph("{{甲}}{{乙}}")
    p.runs[0].bold = True
    doc.add_paragraph("{{丙}}")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "{{表格}}"
    buf = BytesIO()
    doc.save(buf)
    response = client.post("/api/templates", files={"file": ("排版.docx", buf.getvalue())})
    assert response.status_code == 201
    return response.json()


def test_remove_shared_paragraph_and_restore_copy(client: TestClient, layout_scene: dict) -> None:
    tpl = layout_scene
    vid = tpl["default_version_id"]
    selected = next(r for r in tpl["regions"] if r["placeholder"] == "{{甲}}")
    source = (settings.templates_dir / tpl["storage_name"]).read_bytes()
    url = f"/api/versions/{vid}/regions/{selected['id']}/layout"
    assert client.post(url, json={"action": "remove"}).status_code == 200
    with get_conn() as conn:
        regions = regions_repo.list_regions(conn, tpl["id"])
    outcome = apply_replacements(source, regions, {}, removed_region_ids={selected["id"]})
    doc = Document(BytesIO(outcome.data))
    assert [p.text for p in doc.paragraphs] == ["前文", "{{丙}}"]
    hidden = [r for r in regions if r.placeholder in ("{{甲}}", "{{乙}}")]
    assert all(r.id not in outcome.region_paths for r in hidden)
    last = next(r for r in regions if r.placeholder == "{{丙}}")
    assert outcome.region_paths[last.id] == [1]
    copied = client.post(
        f"/api/templates/{tpl['id']}/versions", json={"name": "删除底稿", "copy_from": vid}
    ).json()
    with get_conn() as conn:
        assert (
            conn.execute(
                "SELECT COUNT(*) FROM version_region_actions WHERE version_id=?", (copied["id"],)
            ).fetchone()[0]
            == 1
        )
    other = next(r for r in hidden if r.id != selected["id"])
    assert (
        client.post(
            f"/api/versions/{vid}/regions/{other.id}/layout", json={"action": "restore"}
        ).status_code
        == 200
    )
    with get_conn() as conn:
        assert (
            conn.execute(
                "SELECT COUNT(*) FROM version_region_actions WHERE version_id=?", (vid,)
            ).fetchone()[0]
            == 0
        )
    assert (settings.templates_dir / tpl["storage_name"]).read_bytes() == source


def test_blank_block_keeps_paragraph_and_cell_removal_keeps_valid_cell(
    client: TestClient, layout_scene: dict
) -> None:
    tpl = layout_scene
    block = client.post("/api/blocks", json={"name": "留白", "content": "", "kind": "blank"})
    assert block.status_code == 201 and block.json()["kind"] == "blank"
    assert client.post("/api/blocks", json={"name": "普通", "content": ""}).status_code == 400
    source = (settings.templates_dir / tpl["storage_name"]).read_bytes()
    with get_conn() as conn:
        regions = regions_repo.list_regions(conn, tpl["id"])
    target = next(r for r in regions if r.placeholder == "{{甲}}")
    table = next(r for r in regions if r.placeholder == "{{表格}}")
    outcome = apply_replacements(
        source,
        regions,
        {target.id: ""},
        blank_region_ids={target.id},
        removed_region_ids={table.id},
    )
    doc = Document(BytesIO(outcome.data))
    assert [p.text for p in doc.paragraphs] == ["前文", "", "{{丙}}"]
    assert len(doc.tables[0].cell(0, 0).paragraphs) == 1
    assert doc.tables[0].cell(0, 0).text == ""
    assert target.id in outcome.region_paths and table.id not in outcome.region_paths


def test_layout_rejects_cross_template_and_unknown_action(
    client: TestClient, layout_scene: dict
) -> None:
    tpl = layout_scene
    doc = Document()
    doc.add_paragraph("另一个模板")
    buf = BytesIO()
    doc.save(buf)
    other = client.post("/api/templates", files={"file": ("其他.docx", buf.getvalue())}).json()
    url = f"/api/versions/{other['default_version_id']}/regions/{tpl['regions'][0]['id']}/layout"
    assert client.post(url, json={"action": "remove"}).status_code == 400
    assert client.post(url, json={"action": "invalid"}).status_code == 422


def test_render_removed_blank_and_export(client: TestClient, layout_scene: dict) -> None:
    from app.services.render_service import render_version

    tpl = layout_scene
    vid = tpl["default_version_id"]
    first = next(r for r in tpl["regions"] if r["placeholder"] == "{{甲}}")
    last = next(r for r in tpl["regions"] if r["placeholder"] == "{{丙}}")
    assert (
        client.post(
            f"/api/versions/{vid}/regions/{first['id']}/layout", json={"action": "remove"}
        ).status_code
        == 200
    )
    blank = client.post("/api/blocks", json={"name": "空行", "content": "", "kind": "blank"}).json()
    assert (
        client.post(
            f"/api/versions/{vid}/bindings", json={"region_id": last["id"], "block_id": blank["id"]}
        ).status_code
        == 200
    )
    rendered = render_version(vid)
    removed = next(r for r in rendered.items if r["id"] == first["id"])
    assert removed["removed"] is True and removed["bbox"] is None
    assert [p.text for p in Document(BytesIO(rendered.data)).paragraphs] == ["前文", ""]
    response = client.post(f"/api/versions/{vid}/export", json={"confirm_large_overflow": True})
    assert response.status_code == 200 and response.content == rendered.data


def test_editing_blank_block_can_restore_text(client: TestClient, layout_scene: dict) -> None:
    tpl = layout_scene
    vid = tpl["default_version_id"]
    rid = tpl["regions"][0]["id"]
    block = client.post("/api/blocks", json={"name": "留白", "content": "", "kind": "blank"}).json()
    client.post(f"/api/versions/{vid}/bindings", json={"region_id": rid, "block_id": block["id"]})
    response = client.post(
        f"/api/versions/{vid}/regions/{rid}/text",
        json={"content": "恢复文字", "sync_block": True, "expected_content": ""},
    )
    assert response.status_code == 200
    updated = client.get(f"/api/blocks/{block['id']}").json()
    assert updated["kind"] == "text" and updated["content"] == "恢复文字"
