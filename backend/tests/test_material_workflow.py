"""Data preservation across selective material edits, independent rows and frozen outputs."""

from io import BytesIO
from uuid import uuid4

from docx import Document
from fastapi.testclient import TestClient

from app.models.db import get_conn
from app.services import render_service


def scene(client: TestClient) -> tuple[dict, int, int, int]:
    doc = Document()
    doc.add_paragraph("姓名：{{姓名}}")
    doc.add_paragraph("经历：{{经历}}")
    stream = BytesIO()
    doc.save(stream)
    tpl = client.post("/api/templates", files={"file": ("资料.docx", stream.getvalue())}).json()
    v1 = tpl["default_version_id"]
    v2 = client.post(f"/api/templates/{tpl['id']}/versions", json={"name": "另一稿"}).json()["id"]
    block = client.post("/api/blocks", json={"name": "常用文字", "content": "原素材正文"}).json()[
        "id"
    ]
    for version in (v1, v2):
        for region in tpl["regions"][:2]:
            assert (
                client.post(
                    f"/api/versions/{version}/bindings",
                    json={"region_id": region["id"], "block_id": block},
                ).status_code
                == 200
            )
    return tpl, v1, v2, block


def test_selective_update_preserves_unselected_and_never_matches_copies(client: TestClient) -> None:
    tpl, v1, v2, block = scene(client)
    rid = tpl["regions"][0]["id"]
    # A user explicitly makes a private copy, even with exactly equal text.
    copied = client.post(
        f"/api/blocks/{block}/versions/{v2}/copy", json={"content": "原素材正文"}
    ).json()["id"]
    scope = client.get(f"/api/blocks/{block}/references").json()
    assert {f["version_id"] for f in scope["fields"]} == {v1}
    selected = [{"kind": "region", "version_id": v1, "id": rid}]
    result = client.post(
        f"/api/blocks/{block}/scoped",
        json={
            "token": scope["token"],
            "selected": selected,
            "content": "已同步正文",
        },
    )
    assert result.status_code == 200, result.text
    with get_conn() as conn:
        texts = conn.execute(
            "SELECT b.version_id,b.region_id,b.block_id,k.content FROM bindings b JOIN "
            "blocks k ON k.id=b.block_id ORDER BY b.version_id,b.region_id"
        ).fetchall()
        assert texts[0]["content"] == "已同步正文"
        assert texts[1]["content"] == "原素材正文" and texts[1]["block_id"] != block
        assert all(
            r["block_id"] == copied and r["content"] == "原素材正文"
            for r in texts
            if r["version_id"] == v2
        )
    # Stale selection cannot accidentally affect a newer set of references.
    assert (
        client.post(
            f"/api/blocks/{block}/scoped", json={"token": scope["token"], "content": "迟到"}
        ).status_code
        == 409
    )
    assert client.post("/api/history/undo").status_code == 200
    assert client.get(f"/api/blocks/{block}").json()["content"] == "原素材正文"


def test_metadata_only_does_not_detach_and_bad_fields_roll_back(client: TestClient) -> None:
    _, v1, _, block = scene(client)
    scope = client.get(f"/api/blocks/{block}/references").json()
    bad = client.post(
        f"/api/blocks/{block}/scoped",
        json={
            "token": scope["token"],
            "content": "新",
            "selected": [{"kind": "region", "version_id": v1, "id": 999999}],
        },
    )
    assert bad.status_code == 400
    assert client.get(f"/api/blocks/{block}").json()["content"] == "原素材正文"
    good = client.post(
        f"/api/blocks/{block}/scoped",
        json={"token": scope["token"], "name": "新名称", "tags": ["常用"]},
    )
    assert good.status_code == 200
    assert len(client.get(f"/api/blocks/{block}/references").json()["fields"]) == len(
        scope["fields"]
    )


def test_independent_rows_order_copy_undo_migration_and_output(client: TestClient) -> None:
    tpl, v1, _, block = scene(client)
    rid = tpl["regions"][0]["id"]
    before = client.post(
        f"/api/versions/{v1}/insertions", json={"id": rid, "side": "before"}
    ).json()["id"]
    after = client.post(f"/api/versions/{v1}/insertions", json={"id": rid, "side": "after"}).json()[
        "id"
    ]
    middle = client.post(
        f"/api/versions/{v1}/insertions", json={"kind": "insertion", "id": before, "side": "after"}
    ).json()["id"]
    for identity, content in (
        (before, "上方独立文字"),
        (middle, "中间独立文字"),
        (after, "下方独立文字"),
    ):
        response = client.put(
            f"/api/versions/{v1}/insertions/{identity}",
            json={"content": content, "expected_content": ""},
        )
        assert response.status_code == 200, response.text
    rendered = render_service.render_version(v1)
    refs = [item["field_ref"] for item in rendered.items if isinstance(item.get("field_ref"), dict)]
    assert {ref["id"] for ref in refs if isinstance(ref, dict) and ref["kind"] == "insertion"} == {
        before,
        middle,
        after,
    }
    text = [p.text for p in Document(BytesIO(rendered.data)).paragraphs]
    assert text[:4] == ["上方独立文字", "中间独立文字", "姓名：原素材正文", "下方独立文字"]
    assert (
        client.post(
            f"/api/versions/{v1}/regions/{rid}/layout", json={"action": "remove"}
        ).status_code
        == 409
    )
    assert client.delete(f"/api/regions/{rid}").status_code == 409
    clone = client.post(
        f"/api/templates/{tpl['id']}/versions", json={"name": "复制含行", "copy_from": v1}
    ).json()["id"]
    assert [
        p.text for p in Document(BytesIO(render_service.render_version(clone).data)).paragraphs
    ] == text
    assert client.delete(f"/api/versions/{v1}/insertions/{middle}").status_code == 204
    assert client.post("/api/history/undo").status_code == 200
    assert [
        p.text for p in Document(BytesIO(render_service.render_version(v1).data)).paragraphs
    ] == text
    assert client.post("/api/history/redo").status_code == 200
    newer = client.post(f"/api/versions/{v1}/insertions", json={"id": rid, "side": "after"}).json()[
        "id"
    ]
    assert newer > middle
    # Empty row persists as a legal independent paragraph.
    assert "" in [
        p.text for p in Document(BytesIO(render_service.render_version(v1).data)).paragraphs
    ]
    target_doc = Document()
    target_doc.add_paragraph("新模板 {{姓名}}")
    target_data = BytesIO()
    target_doc.save(target_data)
    target = client.post(
        "/api/templates", files={"file": ("新模板.docx", target_data.getvalue())}
    ).json()
    plan = client.post(
        f"/api/templates/{target['id']}/migrate/plan",
        json={"source_version_id": v1, "mode": "new_document"},
    ).json()
    sources = {item["source_region_id"]: item for item in plan["source_items"]}
    assert sources[-before]["content"] == "上方独立文字"
    assert sources[-after]["content"] == "下方独立文字"
    assert (
        client.post(
            f"/api/templates/{target['id']}/migrate/apply",
            json={
                "source_version_id": v1,
                "mode": "new_document",
                "name": "带入插入文字",
                "plan_token": plan["plan_token"],
                "idempotency_key": uuid4().hex,
                "placements": [
                    {"source_region_id": -before, "target_region_id": target["regions"][0]["id"]}
                ],
                "omitted_source_region_ids": [
                    identity for identity in sources if identity != -before
                ],
            },
        ).status_code
        == 201
    )
    # Selective deletion blanks one row; every unselected material reference keeps its text.
    assert (
        client.put(f"/api/versions/{v1}/insertions/{after}", json={"block_id": block}).status_code
        == 200
    )
    scope = client.get(f"/api/blocks/{block}/references").json()
    assert (
        client.post(
            f"/api/blocks/{block}/scoped",
            json={
                "action": "delete",
                "token": scope["token"],
                "selected": [
                    {"kind": "insertion", "version_id": v1, "id": after},
                    {"kind": "region", "version_id": v1, "id": rid},
                ],
            },
        ).status_code
        == 200
    )
    blank = render_service.render_version(v1)
    paragraphs = [p.text for p in Document(BytesIO(blank.data)).paragraphs]
    assert "姓名：" in paragraphs and "经历：原素材正文" in paragraphs
    assert blank.items[0]["current_text"] == ""
    for endpoint, extra in (("export", {}), ("archives", {"name": "留空归档"})):
        blocked = client.post(f"/api/versions/{v1}/{endpoint}", json=extra)
        assert blocked.status_code == 409, blocked.text
        assert len([w for w in blocked.json()["warnings"] if w.get("kind") == "missing"]) == 2
        assert (
            client.post(
                f"/api/versions/{v1}/{endpoint}", json={**extra, "confirm_large_overflow": True}
            ).status_code
            == 409
        )
        token = blocked.json()["warning_token"]
        confirmed = client.post(
            f"/api/versions/{v1}/{endpoint}",
            json={**extra, "confirm_large_overflow": True, "confirm_warnings_token": token},
        )
        assert confirmed.status_code in (200, 201), confirmed.text
    with get_conn() as conn:
        snapshot = conn.execute("SELECT source_json FROM archives").fetchone()[0]
        assert "missing" in snapshot


def test_set_as_material_is_atomic_and_blank_row_has_own_binding(client: TestClient) -> None:
    tpl, v1, v2, _ = scene(client)
    rid = tpl["regions"][0]["id"]
    with get_conn() as conn:
        conn.execute(
            "UPDATE bindings SET line_break_mode='soft',position='before' "
            "WHERE version_id=? AND region_id=?",
            (v1, rid),
        )
    payload = {
        "name": "新素材",
        "content": "编辑中的内容",
        "tags": ["可复用"],
        "version_id": v1,
        "field": {"kind": "region", "id": rid, "version_id": v1},
        "expected_content": "过期文字",
        "idempotency_key": uuid4().hex,
    }
    with get_conn() as conn:
        before = conn.execute("SELECT COUNT(*) FROM blocks").fetchone()[0]
    assert client.post("/api/blocks/from-field", json=payload).status_code == 409
    with get_conn() as conn:
        assert conn.execute("SELECT COUNT(*) FROM blocks").fetchone()[0] == before
    payload["expected_content"] = "原素材正文"
    created = client.post("/api/blocks/from-field", json=payload)
    assert created.status_code == 201, created.text
    with get_conn() as conn:
        layout = conn.execute(
            "SELECT line_break_mode,position FROM bindings WHERE version_id=? AND region_id=?",
            (v1, rid),
        ).fetchone()
        assert tuple(layout) == ("soft", "before")
    assert (
        client.get(f"/api/versions/{v1}/bindings").json()["bindings"][0]["block_id"]
        == created.json()["id"]
    )
    assert (
        client.get(f"/api/versions/{v2}/bindings").json()["bindings"][0]["block_id"]
        != created.json()["id"]
    )

    replay = client.post("/api/blocks/from-field", json=payload)
    assert replay.status_code == 201 and replay.json()["id"] == created.json()["id"]
    with get_conn() as conn:
        assert conn.execute("SELECT COUNT(*) FROM blocks").fetchone()[0] == before + 1
    assert (
        client.post("/api/blocks/from-field", json={**payload, "name": "不同请求"}).status_code
        == 409
    )
    assert client.post("/api/history/undo").status_code == 200
    assert client.post("/api/blocks/from-field", json=payload).status_code == 410


def test_inserted_rows_move_unbound_manual_overlay_without_changing_template_frame(
    client: TestClient,
) -> None:
    import json

    tpl, _, _, _ = scene(client)
    vid = client.post(f"/api/templates/{tpl['id']}/versions", json={"name": "人工框空白稿"}).json()[
        "id"
    ]
    original = render_service.render_version(vid)
    rid = tpl["regions"][0]["id"]
    old_box = next(item["bbox"] for item in original.items if item["id"] == rid)
    assert isinstance(old_box, dict)
    with get_conn() as conn:
        conn.execute(
            "UPDATE regions SET bbox_source='manual',bbox_json=? WHERE id=?",
            (json.dumps(old_box), rid),
        )
    row = client.post(f"/api/versions/{vid}/insertions", json={"id": rid, "side": "before"}).json()[
        "id"
    ]
    client.put(f"/api/versions/{vid}/insertions/{row}", json={"content": "A new independent line"})
    rendered = render_service.render_version(vid)
    moved_box = next(item["bbox"] for item in rendered.items if item["id"] == rid)
    assert isinstance(moved_box, dict)
    assert float(str(moved_box["y0"])) > float(str(old_box["y0"]))
    with get_conn() as conn:
        saved = conn.execute(
            "SELECT bbox_source,bbox_json FROM regions WHERE id=?", (rid,)
        ).fetchone()
        assert saved["bbox_source"] == "manual" and json.loads(saved["bbox_json"]) == old_box
