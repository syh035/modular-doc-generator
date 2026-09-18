"""M5b 校对测试：框选反解（假 LO）+ 状态机 + 自动 ready + P21 bbox 生命周期 + 真实 LO 端到端。"""

import json
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pymupdf
import pytest
from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

from app.models.db import get_conn
from app.models.repositories import regions as regions_repo
from app.services.libreoffice import find_soffice
from app.services.pdf_geometry import RegionGeometry
from app.services.proofread_service import _frame_hits

PARA_A = "第一段候选占位内容用于校对框选测试，需要超过三十个字符的正文段落才会被识别为候选区域"
PARA_B = "第二段候选占位内容乙"


def make_docx(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_docx_eastAsia(*paragraphs: str) -> bytes:
    """显式 eastAsia 字体（M4 事实：python-docx 默认 run 无 rFonts）。"""
    doc = Document()
    for text in paragraphs:
        p = doc.add_paragraph(text)
        for run in p.runs:
            run.font.name = "宋体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def upload(client: TestClient, data: bytes) -> dict:
    resp = client.post(
        "/api/templates",
        files={"file": ("模板.docx", data, "application/octet-stream")},
    )
    assert resp.status_code == 201
    return resp.json()


def patch_fake_convert(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, texts: list[str]
) -> None:
    """假 LO：按段落文本逐行排版出 PDF（与文档流对齐）。"""

    def convert(path: Path) -> Path:
        pdf = tmp_path / f"fake_{abs(hash(tuple(texts)))}.pdf"
        doc = pymupdf.open()
        page = doc.new_page()
        y = 100.0
        for text in texts:
            page.insert_text((72, y), text, fontname="china-s")
            y += 40.0
        doc.save(pdf)
        doc.close()
        return pdf

    monkeypatch.setattr("app.services.libreoffice._manager", SimpleNamespace(convert=convert))


def static_fake_convert(monkeypatch: pytest.MonkeyPatch, pdf_path: Path) -> None:
    """假 LO：恒返回同一份 PDF（模拟缓存命中 / 产物切换）。"""
    monkeypatch.setattr(
        "app.services.libreoffice._manager", SimpleNamespace(convert=lambda p: pdf_path)
    )


def frame(page: float, x0: float, y0: float, x1: float, y1: float) -> dict:
    return {"page": page, "x0": x0, "y0": y0, "x1": x1, "y1": y1}


def post_region(client: TestClient, tpl_id: int, label: str, type_: str, bbox: dict):
    return client.post(
        f"/api/templates/{tpl_id}/regions",
        json={"label": label, "type": type_, "bbox": bbox},
    )


def make_pdf(tmp_path: Path, texts: list[str], name: str) -> Path:
    doc = pymupdf.open()
    page = doc.new_page()
    y = 100.0
    for text in texts:
        page.insert_text((72, y), text, fontname="china-s")
        y += 40.0
    pdf = tmp_path / name
    doc.save(pdf)
    doc.close()
    return pdf


# ---- 纯函数：_frame_hits 命中判定 ----


def geo(page: int, bbox: tuple[float, float, float, float]) -> RegionGeometry:
    return RegionGeometry(page=page, bbox=bbox, line_count=1)


GEOMETRY = {
    (0,): geo(0, (72.0, 90.0, 300.0, 105.0)),
    (1,): geo(0, (72.0, 130.0, 300.0, 145.0)),
    (2,): geo(1, (72.0, 90.0, 300.0, 105.0)),  # 第 2 页
}


def test_frame_hits_single_and_sliver() -> None:
    # 完整覆盖第 1 段 → 命中 1 个
    assert [p for p, _ in _frame_hits(GEOMETRY, 0, (70.0, 88.0, 310.0, 107.0))] == [(0,)]
    # 发丝级交叠（高 1pt < 2pt）→ 不命中
    assert _frame_hits(GEOMETRY, 0, (70.0, 104.0, 310.0, 129.0)) == []
    # 第 2 页段落（geo.page=1）的命中：page=1 查询命中 (2,)
    assert [p for p, _ in _frame_hits(GEOMETRY, 1, (70.0, 88.0, 310.0, 107.0))] == [(2,)]
    # 覆盖两段 → 命中 2 个
    assert len(_frame_hits(GEOMETRY, 0, (70.0, 88.0, 310.0, 147.0))) == 2


# ---- 框选新建（假 LO 反解）----


def test_create_manual_region_resolves_anchor(client, tmp_path, monkeypatch) -> None:
    body = upload(client, make_docx(PARA_A))
    patch_fake_convert(monkeypatch, tmp_path, [PARA_A])

    resp = post_region(
        client, body["id"], "手选段落", "work", frame(0, 70, 85, 320, 108)
    )
    assert resp.status_code == 201
    region = resp.json()
    assert region["label"] == "手选段落"
    assert region["type"] == "work"
    assert region["review_status"] == "confirmed"  # PRD：完成录入即已确认
    assert region["anchor"] == {"kind": "p", "path": [0]}  # 反解回文档流段落
    assert region["bbox"] == {"page": 0, "x0": 70.0, "y0": 85.0, "x1": 320.0, "y1": 108.0}
    assert region["confidence"] is None
    # 落库事实：manual 来源 + bbox_pdf_sha = 产物 PDF sha
    with get_conn() as conn:
        row = regions_repo.get_region(conn, region["id"])
    assert row is not None and row.bbox_source == "manual"
    assert row is not None and row.bbox_pdf_sha is not None
    # order_index 接在文档流序尾部
    assert region["order_index"] == body["regions"][0]["order_index"] + 1


def test_create_manual_region_empty_frame_rejected(client, tmp_path, monkeypatch) -> None:
    body = upload(client, make_docx(PARA_A))
    patch_fake_convert(monkeypatch, tmp_path, [PARA_A])

    resp = post_region(client, body["id"], "空框", "custom", frame(0, 5, 5, 30, 20))
    assert resp.status_code == 400
    assert resp.json()["error"]["code"] == "REGION_FRAME_EMPTY"


def test_create_manual_region_multi_paragraph_rejected(client, tmp_path, monkeypatch) -> None:
    body = upload(client, make_docx_eastAsia(PARA_A, PARA_B))
    patch_fake_convert(monkeypatch, tmp_path, [PARA_A, PARA_B])

    resp = post_region(client, body["id"], "跨段框", "work", frame(0, 70, 85, 320, 150))
    assert resp.status_code == 400
    assert resp.json()["error"]["code"] == "REGION_FRAME_MULTI"
    assert "2" in resp.json()["error"]["message"]


def test_create_manual_region_invalid_payload(client, tmp_path, monkeypatch) -> None:
    body = upload(client, make_docx(PARA_A))
    patch_fake_convert(monkeypatch, tmp_path, [PARA_A])

    # 倒置矩形 / 空名称
    r1 = post_region(client, body["id"], "x", "work", frame(0, 100, 90, 90, 110))
    r2 = post_region(client, body["id"], "  ", "work", frame(0, 70, 85, 320, 108))
    assert r1.status_code == 400 and r1.json()["error"]["code"] == "REGION_INVALID"
    assert r2.status_code == 400 and r2.json()["error"]["code"] == "REGION_INVALID"


def test_create_region_template_404(client, monkeypatch) -> None:
    monkeypatch.setattr("app.services.libreoffice._manager", SimpleNamespace(convert=lambda p: p))
    resp = post_region(client, 999, "x", "work", frame(0, 1, 1, 10, 10))
    assert resp.status_code == 404
    assert resp.json()["error"]["code"] == "TEMPLATE_NOT_FOUND"


# ---- 状态机 + 自动 ready ----


def test_review_status_transitions(client) -> None:
    body = upload(client, make_docx("电话：{{手机号}}，姓名：{{姓名}}"))
    rid = body["regions"][0]["id"]
    tpl_id = body["id"]

    # pending → confirmed
    r = client.patch(f"/api/regions/{rid}", json={"review_status": "confirmed"})
    assert r.status_code == 200 and r.json()["review_status"] == "confirmed"
    # confirmed → excluded 非法
    r = client.patch(f"/api/regions/{rid}", json={"review_status": "excluded"})
    assert r.status_code == 400 and r.json()["error"]["code"] == "REGION_REVIEW_INVALID"
    # confirmed → pending（重新校对）→ excluded
    r_pending = client.patch(f"/api/regions/{rid}", json={"review_status": "pending"})
    r_excluded = client.patch(f"/api/regions/{rid}", json={"review_status": "excluded"})
    assert r_pending.status_code == 200 and r_excluded.status_code == 200
    # 同状态重复提交视为无变化
    r_repeat = client.patch(f"/api/regions/{rid}", json={"review_status": "excluded"})
    assert r_repeat.status_code == 200
    detail = client.get(f"/api/templates/{tpl_id}").json()
    assert detail["regions"][0]["review_status"] == "excluded"


def test_auto_ready_when_all_reviewed(client) -> None:
    body = upload(client, make_docx("电话：{{手机号}}，姓名：{{姓名}}"))
    tpl_id = body["id"]
    r1, r2 = (r["id"] for r in body["regions"])

    assert client.get(f"/api/templates/{tpl_id}").json()["status"] == "pending_review"
    client.patch(f"/api/regions/{r1}", json={"review_status": "confirmed"})
    assert client.get(f"/api/templates/{tpl_id}").json()["status"] == "pending_review"
    client.patch(f"/api/regions/{r2}", json={"review_status": "excluded"})
    assert client.get(f"/api/templates/{tpl_id}").json()["status"] == "ready"  # 全部处理完毕


def test_auto_ready_after_deleting_last_pending(client) -> None:
    body = upload(client, make_docx("电话：{{手机号}}，姓名：{{姓名}}"))
    tpl_id = body["id"]
    r1, r2 = (r["id"] for r in body["regions"])
    client.patch(f"/api/regions/{r1}", json={"review_status": "confirmed"})
    assert client.delete(f"/api/regions/{r2}").status_code == 204
    assert client.get(f"/api/templates/{tpl_id}").json()["status"] == "ready"
    # 删除后区域消失；重复删除 404
    assert client.delete(f"/api/regions/{r2}").status_code == 404
    remaining = client.get(f"/api/templates/{tpl_id}").json()["regions"]
    assert [r["id"] for r in remaining] == [r1]  # 仅被删的 r2 消失


def test_excluded_region_binding_rejected(client) -> None:
    body = upload(client, make_docx("{{姓名}}"))
    rid = body["regions"][0]["id"]
    vid = body["default_version_id"]
    block = client.post("/api/blocks", json={"name": "姓名块", "content": "张三"})
    client.patch(f"/api/regions/{rid}", json={"review_status": "excluded"})
    resp = client.post(
        f"/api/versions/{vid}/bindings",
        json={"region_id": rid, "block_id": block.json()["id"]},
    )
    assert resp.status_code == 400
    assert resp.json()["error"]["code"] == "REGION_EXCLUDED"


def test_delete_region_cascades_binding(client) -> None:
    body = upload(client, make_docx("{{姓名}}"))
    rid = body["regions"][0]["id"]
    vid = body["default_version_id"]
    block = client.post("/api/blocks", json={"name": "姓名块", "content": "张三"}).json()["id"]
    bound = client.post(
        f"/api/versions/{vid}/bindings", json={"region_id": rid, "block_id": block}
    )
    assert bound.status_code == 200
    client.delete(f"/api/regions/{rid}")
    bindings = client.get(f"/api/versions/{vid}/bindings").json()["bindings"]
    assert bindings == []


# ---- P21 bbox 生命周期：auto 随 PDF 失效重算，manual 保护 ----


def test_bbox_lifecycle_auto_recompute_manual_protected(client, tmp_path, monkeypatch) -> None:
    """P21：auto bbox 随渲染产物变化失效重算（sha 跟随）；manual 校对产物保护不覆盖。"""
    body = upload(client, make_docx(PARA_A))
    tpl_id = body["id"]
    rid = body["regions"][0]["id"]

    # 首次预览：auto bbox 落库（含来源 sha）
    patch_fake_convert(monkeypatch, tmp_path, [PARA_A])
    assert client.get(f"/api/templates/{tpl_id}/preview").status_code == 200
    with get_conn() as conn:
        row = regions_repo.get_region(conn, rid)
    assert row is not None and row.bbox_json is not None and row.bbox_source == "auto"
    sha1 = row.bbox_pdf_sha
    assert sha1 is not None

    # 渲染产物变化（追加干扰行 → PDF 内容不同 → sha 变）→ auto bbox 失效重算
    alt = make_pdf(tmp_path, [PARA_A, "干扰行"], "alt.pdf")
    static_fake_convert(monkeypatch, alt)
    assert client.get(f"/api/templates/{tpl_id}/preview").status_code == 200
    with get_conn() as conn:
        row2 = regions_repo.get_region(conn, rid)
    assert row2 is not None and row2.bbox_pdf_sha is not None and row2.bbox_pdf_sha != sha1

    # 人工微调 bbox → manual；产物再变，manual bbox 与来源 sha 保持不动
    client.patch(
        f"/api/regions/{rid}",
        json={"bbox": {"page": 0, "x0": 10, "y0": 10, "x1": 60, "y1": 30}},
    )
    alt2 = make_pdf(tmp_path, ["完全不同的产物"], "alt2.pdf")
    static_fake_convert(monkeypatch, alt2)
    assert client.get(f"/api/templates/{tpl_id}/preview").status_code == 200
    with get_conn() as conn:
        row3 = regions_repo.get_region(conn, rid)
    assert row3 is not None
    assert row3.bbox_source == "manual"
    assert json.loads(row3.bbox_json or "{}") == {
        "page": 0, "x0": 10.0, "y0": 10.0, "x1": 60.0, "y1": 30.0,
    }
    assert row3.bbox_pdf_sha == row2.bbox_pdf_sha


# ---- 真实 LO 端到端 ----


@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice 未安装")
def test_manual_region_real_lo(client) -> None:
    """框选矩形 = 既有区域 bbox → 反解回同一段落；状态机走通至 ready。"""
    data = make_docx_eastAsia("工作经历：负责核心系统研发", "项目经历：主导数据平台建设")
    body = upload(client, data)
    tpl_id = body["id"]

    assert client.get(f"/api/templates/{tpl_id}/preview").status_code == 200
    regions = client.get(f"/api/templates/{tpl_id}/regions").json()["regions"]
    target = regions[0]
    assert target["bbox"] is not None

    resp = client.post(
        f"/api/templates/{tpl_id}/regions",
        json={"label": "手选工作经历", "type": "work", "bbox": target["bbox"]},
    )
    assert resp.status_code == 201
    created = resp.json()
    assert created["anchor"]["path"] == target["anchor"]["path"]  # 反解回同一段落
    assert created["bbox"] == target["bbox"]

    # 逐个确认 → 自动 ready
    client.patch(f"/api/regions/{created['id']}", json={"review_status": "confirmed"})
    for r in regions:
        client.patch(f"/api/regions/{r['id']}", json={"review_status": "confirmed"})
    assert client.get(f"/api/templates/{tpl_id}").json()["status"] == "ready"
