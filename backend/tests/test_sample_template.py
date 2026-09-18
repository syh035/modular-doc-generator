"""示例模板服务 + /api/guide/sample-template 端点测试（UI 调整①配套）。

P17 教训验证点：必须逐项检查 rFonts eastAsia 声明存在；
文本可提取 ≠ 字形渲染正常，字体声明是渲染不空白的防线。
"""

import io

from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

from app.services.sample_template import build_sample_docx

_EXPECTED_PLACEHOLDERS = ["{{姓名}}", "{{手机号}}", "{{教育经历}}", "{{工作经历}}", "{{自我评价}}"]


def test_build_sample_docx_reopenable_with_placeholders() -> None:
    data = build_sample_docx()
    assert data[:2] == b"PK"  # zip 魔数
    doc = Document(io.BytesIO(data))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    for ph in _EXPECTED_PLACEHOLDERS:
        assert ph in full_text


def test_build_sample_docx_declares_east_asia_font() -> None:
    doc = Document(io.BytesIO(build_sample_docx()))
    runs = [run for p in doc.paragraphs for run in p.runs]
    assert runs  # 至少有 run
    for run in runs:
        rpr = run._element.rPr  # noqa: SLF001
        assert rpr is not None
        rfonts = rpr.rFonts
        assert rfonts is not None
        assert rfonts.get(qn("w:eastAsia")) is not None  # 中文字形声明防线（P17）


def test_get_sample_template_endpoint(client: TestClient) -> None:
    resp = client.get("/api/guide/sample-template")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "attachment" in resp.headers["content-disposition"]
    doc = Document(io.BytesIO(resp.content))
    assert "{{姓名}}" in "\n".join(p.text for p in doc.paragraphs)
