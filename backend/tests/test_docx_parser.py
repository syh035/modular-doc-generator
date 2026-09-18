"""docx_parser 单测：P1 run 合并、表格/嵌套表格、D6 页眉页脚与文本框跳过。"""

from io import BytesIO

from docx import Document
from lxml import etree

from app.services.docx_parser import iter_flow_paragraphs, parse_candidates, parse_placeholders

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_V_NS = "urn:schemas-microsoft-com:vml"


def docx_bytes(doc: Document) -> bytes:
    """Document → bytes（内存序列化）。"""
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---- 构造辅助 ----


def build_simple_doc() -> Document:
    doc = Document()
    doc.add_paragraph("{{姓名}}")
    return doc


def build_split_run_doc() -> Document:
    """P1 核心场景：`{{姓名}}` 被拆成多个 run。"""
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("{{姓")
    p.add_run("名")
    p.add_run("}}")
    return doc


def build_multi_placeholder_doc() -> Document:
    doc = Document()
    doc.add_paragraph("{{姓名}}，电话{{电话}}")
    doc.add_paragraph("{{ 自我评价 }}")
    return doc


def build_table_doc() -> Document:
    doc = Document()
    doc.add_paragraph("正文占位 {{姓名}}")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "{{电话}}"
    table.cell(1, 1).text = "{{邮箱}}"
    return doc


def build_nested_table_doc() -> Document:
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = "外层"
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "{{嵌套}}"
    return doc


def build_textbox_doc() -> Document:
    """D6：文本框内占位符不纳入解析范围。"""
    doc = Document()
    p = doc.add_paragraph()
    pict_run = etree.fromstring(
        f'<w:r xmlns:w="{_W_NS}">'
        f'<w:pict><v:shape xmlns:v="{_V_NS}">'
        "<v:textbox><w:txbxContent>"
        "<w:p><w:r><w:t>{{文本框占位}}</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict></w:r>"
    )
    p._p.append(pict_run)
    doc.add_paragraph("{{正常区域}}")
    return doc


def build_header_doc() -> Document:
    """D6：页眉中的占位符不纳入解析范围（页眉在独立 header part）。"""
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "{{页眉占位}}"
    doc.add_paragraph("{{正文区域}}")
    return doc


# ---- 测试 ----


def test_simple_placeholder() -> None:
    regions = parse_placeholders(docx_bytes(build_simple_doc()))
    assert len(regions) == 1
    r = regions[0]
    assert r.placeholder == "{{姓名}}"
    assert r.label == "姓名"
    assert r.anchor["kind"] == "p"
    assert len(r.anchor["path"]) == 1  # type: ignore[arg-type]
    assert r.order_index == 0


def test_split_run_merged() -> None:
    """P1：跨 run 拆分的占位符必须能合并识别。"""
    regions = parse_placeholders(docx_bytes(build_split_run_doc()))
    assert len(regions) == 1
    assert regions[0].label == "姓名"


def test_multiple_placeholders_and_trim() -> None:
    regions = parse_placeholders(docx_bytes(build_multi_placeholder_doc()))
    assert [r.label for r in regions] == ["姓名", "电话", "自我评价"]
    assert [r.order_index for r in regions] == [0, 1, 2]


def test_table_cells_parsed() -> None:
    regions = parse_placeholders(docx_bytes(build_table_doc()))
    labels = [r.label for r in regions]
    assert labels == ["姓名", "电话", "邮箱"]
    # 表格区域锚点：[tbl, row, cell, para]，row/cell/para 从 0 起
    phone, mail = regions[1], regions[2]
    assert phone.anchor["kind"] == "cell_p"
    assert phone.anchor["path"] == [phone.anchor["path"][0], 0, 0, 0]  # type: ignore[index]
    assert mail.anchor["path"] == [mail.anchor["path"][0], 1, 1, 0]  # type: ignore[index]


def test_nested_table_parsed() -> None:
    regions = parse_placeholders(docx_bytes(build_nested_table_doc()))
    assert len(regions) == 1
    r = regions[0]
    assert r.label == "嵌套"
    # 嵌套路径：[tbl, row, cell, nested_tbl, row, cell, para] 共 7 位
    path = r.anchor["path"]
    assert isinstance(path, list) and len(path) == 7  # type: ignore[arg-type]


def test_textbox_skipped() -> None:
    """D6：文本框内占位符不识别，只识别正文段落。"""
    regions = parse_placeholders(docx_bytes(build_textbox_doc()))
    assert [r.label for r in regions] == ["正常区域"]


def test_header_skipped() -> None:
    """D6：页眉占位符不识别，只识别正文段落。"""
    regions = parse_placeholders(docx_bytes(build_header_doc()))
    assert [r.label for r in regions] == ["正文区域"]


def test_no_placeholder_returns_empty() -> None:
    doc = Document()
    doc.add_paragraph("普通文本段落，没有占位符")
    doc.add_paragraph("另一个 { 单花括号 } 也不是")
    assert parse_placeholders(docx_bytes(doc)) == []


def test_malformed_braces_not_matched() -> None:
    doc = Document()
    doc.add_paragraph("{{{三重}}}、{{}}、}{反向} 都不算占位符")
    regions = parse_placeholders(docx_bytes(doc))
    # "{{{三重}}}" 中 "{{三重}}" 会被宽匹配到（畸形输入的容错行为，可接受），
    # 但 "{{}}"（空名）与反向必须不产出
    labels = [r.label for r in regions]
    assert "" not in labels
    assert "}" not in labels


# ---- iter_flow_paragraphs（M4 匹配数据源，与解析同源）----


def test_flow_enumerates_all_paragraphs_in_order() -> None:
    """文档流全序枚举：无占位符段落也在列，顺序 = 解析顺序。"""
    doc = Document()
    doc.add_paragraph("标题")
    doc.add_paragraph("电话：{{手机号}}")
    doc.add_paragraph("")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "{{教育}}"
    table.cell(0, 1).text = "格内普通段"
    doc.add_paragraph("结尾")

    flow = list(iter_flow_paragraphs(docx_bytes(doc)))
    texts = [t for _, t in flow]
    assert texts == ["标题", "电话：{{手机号}}", "", "{{教育}}", "格内普通段", "结尾"]
    kinds = [a["kind"] for a, _ in flow]
    assert kinds == ["p", "p", "p", "cell_p", "cell_p", "p"]


def test_flow_anchors_match_parse_placeholders() -> None:
    """锚点一致性：占位符区域的 anchor 必能在文档流枚举中找到同 path 同文本。"""
    doc = build_table_doc()
    data = docx_bytes(doc)
    flow = list(iter_flow_paragraphs(data))
    for region in parse_placeholders(data):
        path = region.anchor["path"]
        matches = [
            t
            for a, t in flow
            if a["path"] == path and a["kind"] == region.anchor["kind"]
        ]
        assert len(matches) == 1
        assert region.placeholder in matches[0]


def test_flow_nested_table_paths() -> None:
    """嵌套表格路径编码：外层段 → 嵌套段 → 尾空段（OOXML 要求 tc 以 p 结尾，
    python-docx 的 add_table 会自动补尾部空段，占外层 cell 的 para_idx=1）。"""
    flow = list(iter_flow_paragraphs(docx_bytes(build_nested_table_doc())))
    assert [a["kind"] for a, _ in flow] == ["cell_p", "cell_p", "cell_p"]
    assert [t for _, t in flow] == ["外层", "{{嵌套}}", ""]
    nested_path = flow[1][0]["path"]
    assert isinstance(nested_path, list)
    assert len(nested_path) == 7  # [tbl,row,cell, nested_tbl,row,cell, para]
    assert flow[2][0]["path"] == [0, 0, 0, 1]  # 尾空段：外层 cell 第 2 个段落


# ---- parse_candidates（M3b 三层识别）----


def test_candidates_placeholder_type_inference() -> None:
    """决策 C：占位符 label 过词表推断类型，confidence=1.0。"""
    doc = Document()
    doc.add_paragraph("{{姓名}}")
    doc.add_paragraph("{{自定义字段}}")
    outcome = parse_candidates(docx_bytes(doc))
    types = [c.region_type for c in outcome.candidates]
    assert types == ["name", "custom"]
    assert all(c.confidence == 1.0 for c in outcome.candidates)
    assert all(c.placeholder is not None for c in outcome.candidates)


def test_candidates_placeholder_paragraph_excludes_paragraph_candidate() -> None:
    """占位符段不再产出段落级候选（一段不重复建区域）。"""
    doc = Document()
    doc.add_paragraph("工作经历：{{工作经历}}")  # 同时命中词表前缀与占位符
    outcome = parse_candidates(docx_bytes(doc))
    assert len(outcome.candidates) == 1
    assert outcome.candidates[0].placeholder == "{{工作经历}}"


def test_candidates_lexicon_title_and_labeled() -> None:
    doc = Document()
    doc.add_paragraph("教育背景")
    doc.add_paragraph("姓名：李四")
    outcome = parse_candidates(docx_bytes(doc))
    edu, name = outcome.candidates
    assert (edu.region_type, edu.confidence, edu.label) == ("education", 0.9, "教育背景")
    assert edu.placeholder is None
    assert (name.region_type, name.confidence, name.label) == ("name", 0.7, "姓名")
    assert name.order_index == 1


def test_candidates_paragraph_threshold_boundary() -> None:
    """成段阈值：30 字产出、29 字不产出。"""
    doc = Document()
    doc.add_paragraph("字" * 29)
    doc.add_paragraph("字" * 30)
    outcome = parse_candidates(docx_bytes(doc))
    assert len(outcome.candidates) == 1
    c = outcome.candidates[0]
    assert c.region_type == "custom"
    assert c.confidence == 0.4
    assert c.label.endswith("…")


def test_candidates_short_non_lexicon_no_candidate() -> None:
    """过短非词表段（如「张三」）不产候选，但计入 has_any_text。"""
    doc = Document()
    doc.add_paragraph("张三")
    outcome = parse_candidates(docx_bytes(doc))
    assert outcome.candidates == []
    assert outcome.has_any_text is True


def test_candidates_blank_document_not_image_only_signal() -> None:
    """全空文档：无候选且 has_any_text=False（纯图片模板判定依据）。"""
    doc = Document()  # 仅默认空段落
    outcome = parse_candidates(docx_bytes(doc))
    assert outcome.candidates == []
    assert outcome.has_any_text is False


def test_candidates_table_cell_lexicon() -> None:
    """词表识别覆盖表格单元格（D6 扫描范围同占位符）。"""
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "项目经历"
    outcome = parse_candidates(docx_bytes(doc))
    assert len(outcome.candidates) == 1
    assert outcome.candidates[0].region_type == "project"
    assert outcome.candidates[0].anchor["kind"] == "cell_p"
