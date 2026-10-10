"""历史模板一次性全文补齐，保留区域身份、排除、绑定及后续删除。"""

import json
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from app.models.db import get_conn
from app.models.repositories import bindings, blocks, regions
from app.services.template_service import upgrade_full_candidates


def test_once_backfill_preserves_old_regions_and_deleted_candidates(client: TestClient) -> None:
    doc = Document()
    for text in ["短句", "教育背景", "长" * 35, "另一个短句", "{{姓名}}"]:
        doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    tpl = client.post("/api/templates", files={"file": ("历史.docx", buf.getvalue())}).json()
    ids = [r["id"] for r in tpl["regions"]]
    with get_conn() as conn:
        conn.execute("DELETE FROM template_candidate_scans WHERE template_id = ?", (tpl["id"],))
        # 模拟旧规则遗漏短句，同时存在已物理删除的长段落和占位符。
        for i in [0, 2, 3, 4]:
            conn.execute("DELETE FROM regions WHERE id = ?", (ids[i],))
        regions.update_region(
            conn, ids[1], review_status="excluded", bbox={"page": 0, "x0": 1}, bbox_source="manual"
        )
        block = blocks.create_block(conn, "旧块", "旧内容")
        bindings.upsert_binding(conn, tpl["default_version_id"], ids[1], block.id)
    with get_conn() as conn:
        assert upgrade_full_candidates(conn) == 2
        result = regions.list_regions(conn, tpl["id"])
        assert [json.loads(r.anchor)["path"] for r in result] == [[0], [1], [3]]
        assert [r.order_index for r in result] == [0, 1, 2]
        old = regions.get_region(conn, ids[1])
        assert old is not None
        assert old.review_status == "excluded" and old.bbox_source == "manual"
        assert old.bbox_json is not None
        assert json.loads(old.bbox_json) == {"page": 0, "x0": 1}
        found_1 = bindings.get_binding(conn, tpl["default_version_id"], old.id)
        assert found_1 is not None
        assert found_1.block_id == block.id
        conn.execute("DELETE FROM regions WHERE id = ?", (result[0].id,))
    with get_conn() as conn:
        assert upgrade_full_candidates(conn) == 0
        assert len(regions.list_regions(conn, tpl["id"])) == 2


def test_new_upload_registered_and_blank_table_cells_ignored(client: TestClient) -> None:
    doc = Document()
    doc.add_paragraph("姓名短句")
    doc.add_paragraph("  ")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "格子"
    buf = BytesIO()
    doc.save(buf)
    tpl = client.post("/api/templates", files={"file": ("全文.docx", buf.getvalue())}).json()
    assert [r["label"] for r in tpl["regions"]] == ["姓名短句", "格子"]
    with get_conn() as conn:
        assert (
            conn.execute(
                "SELECT revision FROM template_candidate_scans WHERE template_id = ?", (tpl["id"],)
            ).fetchone()[0]
            == 3
        )
        assert upgrade_full_candidates(conn) == 0


def test_corrupt_historical_original_does_not_block_startup(client: TestClient) -> None:
    from app.core.config import settings
    from app.main import create_app

    doc = Document()
    doc.add_paragraph("原文字")
    buf = BytesIO()
    doc.save(buf)
    tpl = client.post("/api/templates", files={"file": ("损坏历史.docx", buf.getvalue())}).json()
    with get_conn() as conn:
        conn.execute("DELETE FROM template_candidate_scans WHERE template_id = ?", (tpl["id"],))
    (settings.templates_dir / tpl["storage_name"]).write_bytes(b"corrupted")
    assert create_app() is not None
    with get_conn() as conn:
        assert len(regions.list_regions(conn, tpl["id"])) == 1
        assert conn.execute("SELECT COUNT(*) FROM template_candidate_scans").fetchone()[0] == 0
