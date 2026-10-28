from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient


def test_review_inventory_does_not_create_a_draft(client: TestClient) -> None:
    doc = Document()
    doc.add_paragraph("{{姓名}}")
    doc.sections[0].header.paragraphs[0].text = "页眉保留"
    doc.sections[0].footer.paragraphs[0].text = "页脚保留"
    data = BytesIO()
    doc.save(data)
    uploaded = client.post(
        "/api/templates",
        files={"file": ("检查.docx", data.getvalue())},
        data={"create_default_version": "false"},
    ).json()
    template_id = uploaded["id"]
    result = client.get(f"/api/templates/{template_id}/review-summary")
    assert result.status_code == 200
    assert {item["kind"]: item["text"] for item in result.json()["unsupported"]} == {
        "页眉": "页眉保留",
        "页脚": "页脚保留",
    }
    assert client.get(f"/api/templates/{template_id}/versions").json()["versions"] == []
