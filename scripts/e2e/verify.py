"""E2E 产物校验（backend/.venv/bin/python 运行，三场景 E2E 专用）。

用法（仓库根目录）：
    backend/.venv/bin/python scripts/e2e/verify.py preview \
        --template-id N [--expect 文本 ...] [--absent 文本 ...]
    backend/.venv/bin/python scripts/e2e/verify.py export-file \
        --path 文件.docx [--expect 文本 ...] [--absent 文本 ...]

preview：模板详情 → default_version_id → GET /versions/{id}/preview →
         后端 pdf_geometry.extract_pdf_lines 提取全文 → 逐条断言（单管线产物，
         与预览同源；必须走视觉行提取，P18 裸 get_text() 字符级乱序）。
export-file：zip 魔数校验 + python-docx 提取全文（正文+表格）→ 逐条断言。
全部通过输出 {"ok": true, ...} 并 exit 0；任一失败 exit 1。
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import tempfile
from pathlib import Path

import httpx

_BASE_DIR = Path(__file__).resolve().parent  # scripts/e2e/
BACKEND = _BASE_DIR.parent.parent / "backend"  # scripts/e2e/ → scripts/ → 仓库根/backend
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

BASE = "http://127.0.0.1:8740"


def _norm(text: str) -> str:
    """空白归一：LO 内容流会把同行文字拆片（P18），提取文本含杂乱空白，
    CJK 语义下空白无意义，匹配前双方剔除。"""
    return re.sub(r"\s+", "", text)


def _assert_text(text: str, expects: list[str], absents: list[str]) -> list[str]:
    problems: list[str] = []
    normalized = _norm(text)
    for needle in expects:
        if _norm(needle) not in normalized:
            problems.append(f"缺失预期文本: {needle!r}")
    for needle in absents:
        if _norm(needle) in normalized:
            problems.append(f"出现禁止文本: {needle!r}")
    return problems


def cmd_preview(args: argparse.Namespace) -> dict:
    from app.services.pdf_geometry import extract_pdf_lines  # noqa: E402

    detail = httpx.get(f"{BASE}/api/templates/{args.template_id}", timeout=30).json()
    version_id = detail.get("default_version_id")
    if version_id is None:
        raise SystemExit(json.dumps({"ok": False, "error": "模板无默认版本"}, ensure_ascii=False))
    resp = httpx.get(f"{BASE}/api/versions/{version_id}/preview", timeout=180)
    resp.raise_for_status()
    # extract_pdf_lines 收文件路径（P18 视觉行提取），响应字节先落临时文件
    fd, tmp_name = tempfile.mkstemp(suffix=".pdf")
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(resp.content)
        lines = extract_pdf_lines(tmp_name)
    finally:
        os.unlink(tmp_name)
    text = "\n".join(ln.text for ln in lines)
    problems = _assert_text(text, args.expect, args.absent)
    return {
        "ok": not problems,
        "version_id": version_id,
        "pages": max((ln.page for ln in lines), default=-1) + 1,
        "problems": problems,
    }


def cmd_export_file(args: argparse.Namespace) -> dict:
    from docx import Document

    path = Path(args.path)
    if not path.is_file():
        return {"ok": False, "problems": [f"文件不存在: {path}"]}
    if path.read_bytes()[:2] != b"PK":
        return {"ok": False, "problems": ["非 zip 魔数，产物不是合法 DOCX"]}
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    problems = _assert_text("\n".join(parts), args.expect, args.absent)
    return {"ok": not problems, "path": str(path), "problems": problems}


def main() -> None:
    parser = argparse.ArgumentParser(description="E2E 产物校验")
    sub = parser.add_subparsers(dest="command", required=True)

    p_preview = sub.add_parser("preview", help="校验版本预览 PDF 文本")
    p_preview.add_argument("--template-id", type=int, required=True)
    _add_text_args(p_preview)

    p_export = sub.add_parser("export-file", help="校验导出 DOCX 文件")
    p_export.add_argument("--path", required=True)
    _add_text_args(p_export)

    args = parser.parse_args()
    result = cmd_preview(args) if args.command == "preview" else cmd_export_file(args)
    print(json.dumps(result, ensure_ascii=False))
    sys.exit(0 if result["ok"] else 1)


def _add_text_args(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--expect", action="append", default=[], help="必须包含的文本")
    parser.add_argument("--absent", action="append", default=[], help="必须不存在的文本")


if __name__ == "__main__":
    main()
