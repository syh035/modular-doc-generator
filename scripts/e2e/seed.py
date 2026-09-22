"""E2E fixture 播种（backend/.venv/bin/python 运行，三场景 E2E 专用）。

用法（仓库根目录）：
    backend/.venv/bin/python scripts/e2e/seed.py <scenario>

scenario ∈ {1, 2, 3}：
  1 日常沉淀源模板：{{姓名}} + {{工作经历}} 占位符
  2 定向投递模板：固定行高表格（hRule=exact，P6 裁剪）+ {{概述}} 占位符
  3 模板换装对：迁移源（姓名/工作经历）+ 迁移目标（同类型、布局对调）

fixture 内嵌时间戳行 → 每次 SHA-256 不同 → 恒为全新模板（pending_review，
校对确认在 E2E 浏览器流内完成），保证重复运行确定性。
成功输出 JSON：{"ok": true, "templates": [{id, filename, status}, ...]}。
"""

from __future__ import annotations

import io
import json
import sys
import time

import httpx
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

BASE = "http://127.0.0.1:8740"
TS = time.strftime("%H%M%S")


def _run(doc: Document, text: str, *, size: int = 12, bold: bool = False) -> None:
    """加段落并显式设置 eastAsia 字体（P17：缺省无中文字形，渲染空白）。"""
    r = doc.add_paragraph().add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def _exact_row_table(doc: Document, text: str, height_twips: int = 600) -> None:
    """单格表格 + 固定行高 hRule=exact（P6：内容超出即被裁剪 → 强制大超出）。"""
    table = doc.add_table(rows=1, cols=1)
    trHeight = OxmlElement("w:trHeight")
    trHeight.set(qn("w:val"), str(height_twips))
    trHeight.set(qn("w:hRule"), "exact")
    table.rows[0]._tr.get_or_add_trPr().append(trHeight)
    r = table.rows[0].cells[0].paragraphs[0].add_run(text)
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def _upload(doc: Document, filename: str) -> dict:
    buf = io.BytesIO()
    doc.save(buf)
    resp = httpx.post(
        f"{BASE}/api/templates",
        files={"file": (filename, buf.getvalue())},
        timeout=60,
    )
    body = resp.json()
    if resp.status_code not in (200, 201):
        raise RuntimeError(f"上传失败 {resp.status_code}: {body}")
    return {"id": body["id"], "filename": body["filename"], "status": body["status"]}


def scenario_1() -> list[dict]:
    doc = Document()
    _run(doc, "个人简历", size=18, bold=True)
    _run(doc, "{{姓名}}")
    _run(doc, "{{工作经历}}")
    _run(doc, f"fixture-ts-{TS}")
    return [_upload(doc, f"e2e_s1_日常沉淀_{TS}.docx")]


def scenario_2() -> list[dict]:
    doc = Document()
    _run(doc, "项目概述模板", size=16, bold=True)
    _exact_row_table(doc, "{{概述}}")
    _run(doc, f"fixture-ts-{TS}")
    return [_upload(doc, f"e2e_s2_定向投递_{TS}.docx")]


def scenario_3() -> list[dict]:
    src = Document()
    _run(src, "个人简历", size=18, bold=True)
    _run(src, "{{姓名}}")
    _run(src, "{{工作经历}}")
    _run(src, f"fixture-ts-{TS}")
    tgt = Document()
    _run(tgt, "求职简历", size=20, bold=True)
    _run(tgt, "{{工作经历}}")
    _run(tgt, "{{姓名}}")
    _run(tgt, f"fixture-ts-{TS}")
    return [
        _upload(src, f"e2e_s3_迁移源_{TS}.docx"),
        _upload(tgt, f"e2e_s3_迁移目标_{TS}.docx"),
    ]


def main() -> None:
    scenario = sys.argv[1] if len(sys.argv) > 1 else ""
    builders = {"1": scenario_1, "2": scenario_2, "3": scenario_3}
    if scenario not in builders:
        raise SystemExit(f"用法: seed.py <1|2|3>，收到: {scenario!r}")
    templates = builders[scenario]()
    print(json.dumps({"ok": True, "templates": templates}, ensure_ascii=False))


if __name__ == "__main__":
    main()
