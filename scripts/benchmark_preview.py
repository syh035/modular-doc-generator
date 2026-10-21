"""Measure actual cold/edited/warm previews and export parity in an isolated library."""

import hashlib
import json
import os
import statistics
import sys
import tempfile
import time
from io import BytesIO
from pathlib import Path
from unittest.mock import patch


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="modudoc-benchmark-") as directory:
        os.environ["MODUDOC_DATA_DIR"] = directory
        sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
        import pymupdf
        from app.main import create_app
        from app.services.libreoffice import LibreOfficeManager, get_manager
        from docx import Document
        from fastapi.testclient import TestClient

        conversion_ms = 0.0
        original_run = LibreOfficeManager._run_soffice

        def timed_run(
            manager: LibreOfficeManager, inputs: list[Path], output: Path
        ) -> list[Path]:
            nonlocal conversion_ms
            started = time.perf_counter()
            try:
                return original_run(manager, inputs, output)
            finally:
                conversion_ms += (time.perf_counter() - started) * 1000

        results = []
        with (
            TestClient(create_app()) as client,
            patch.object(LibreOfficeManager, "_run_soffice", timed_run),
        ):
            for count in (1, 5, 10):
                document = Document()
                document.add_paragraph("姓名：{{姓名}}")
                for page in range(count):
                    if page:
                        document.add_page_break()
                    document.add_paragraph(f"第 {page + 1} 页 · 中文文档性能验收")
                    document.add_paragraph("项目经历与工作成果。保留文档样式、中文字体与分页。")
                source = BytesIO()
                document.save(source)
                started = time.perf_counter()
                response = client.post(
                    "/api/templates", files={"file": (f"{count}页.docx", source.getvalue())}
                )
                response.raise_for_status()
                parse_ms = (time.perf_counter() - started) * 1000
                template = response.json()
                version = template["default_version_id"]
                region = next(r for r in template["regions"] if r["placeholder"] is not None)
                block = client.post(
                    "/api/blocks", json={"name": "测试姓名", "content": "测试用户"}
                ).json()
                client.post(
                    f"/api/versions/{version}/bindings",
                    json={"region_id": region["id"], "block_id": block["id"]},
                ).raise_for_status()
                timings = []
                manager = get_manager()
                for index in range(5):
                    if index == 4:
                        client.put(
                            f"/api/blocks/{block['id']}",
                            json={"content": "改稿用户"},
                        ).raise_for_status()
                    conversions_before = manager.conversions
                    conversion_before = conversion_ms
                    started = time.perf_counter()
                    preview = client.get(f"/api/versions/{version}/preview")
                    preview.raise_for_status()
                    preview_ms = (time.perf_counter() - started) * 1000
                    token = preview.headers["x-render-token"]
                    overlay = client.get(f"/api/versions/{version}/overlay?render_token={token}")
                    overlay.raise_for_status()
                    assert (
                        overlay.json()["pdf_sha256"] == hashlib.sha256(preview.content).hexdigest()
                    )
                    timings.append(
                        {
                            "total_ms": round((time.perf_counter() - started) * 1000, 2),
                            "preview_ms": round(preview_ms, 2),
                            "conversion_ms": round(conversion_ms - conversion_before, 2),
                            "other_render_ms": round(
                                preview_ms - (conversion_ms - conversion_before), 2
                            ),
                            "overlay_ms": round(
                                (time.perf_counter() - started) * 1000 - preview_ms, 2
                            ),
                            "conversions": manager.conversions - conversions_before,
                        }
                    )
                with pymupdf.open(stream=preview.content, filetype="pdf") as pdf:
                    assert len(pdf) == count
                    assert all(any(value < 200 for value in p.get_pixmap().samples) for p in pdf)
                    assert "改稿用户" in "".join(page.get_text() for page in pdf)
                assert timings[0]["conversions"] >= 1
                assert all(value["conversions"] == 0 for value in timings[1:4])
                assert timings[4]["conversions"] == 1
                exported = client.post(
                    f"/api/versions/{version}/export", json={"confirm_large_overflow": True}
                )
                exported.raise_for_status()
                path = Path(directory) / f"export-{count}.docx"
                path.write_bytes(exported.content)
                assert get_manager().convert(path).read_bytes() == preview.content
                results.append(
                    {
                        "pages": count,
                        "parse_ms": round(parse_ms, 2),
                        "cold_preview": timings[0],
                        "edited_preview": timings[4],
                        "warm_preview_median_ms": round(
                            statistics.median(value["total_ms"] for value in timings[1:4]), 2
                        ),
                        "preview_export_equal": True,
                    }
                )
        print(
            json.dumps(
                {"samples": results, "scope": "synthetic Chinese documents; local machine"},
                ensure_ascii=False,
                indent=2,
            )
        )


if __name__ == "__main__":
    main()
